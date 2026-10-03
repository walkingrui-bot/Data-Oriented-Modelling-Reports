"""EXP044 original ZIP read and fixed one-feature fall-history table."""

import csv
import io
import json
import math
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
P36 = ROOT / "STAT-PSYMOE-EXP036-20261002-001"
P37 = ROOT / "STAT-PSYMOE-EXP037-20261002-001"
P38 = ROOT / "STAT-PSYMOE-EXP038-20261002-001"
P43 = ROOT / "STAT-PSYMOE-EXP043-20261002-001"
sys.path.insert(0, str(P37))
import audit_two_members as reader  # noqa: E402
sys.path.insert(0, str(P38))
import audit_cohort as cohort  # noqa: E402
sys.path.insert(0, str(P43))
import audit_english_months as date_source  # noqa: E402


def imu_asymmetry(raw):
    parsed = reader.parse_member(raw)
    if cohort.qualify(parsed):
        raise ValueError("IMU_NOT_NUMERICALLY_QUALIFIED: " + ",".join(cohort.qualify(parsed)))
    by_side = {"left": [], "right": []}
    for row in csv.DictReader(io.StringIO(raw.decode("utf-8-sig"), newline="")):
        side = row["foot"].strip().lower()
        if side not in by_side:
            raise ValueError("UNEXPECTED_FOOT_CODE")
        a = math.sqrt(sum(float(row[name]) ** 2 for name in ("acx", "acy", "acz")))
        if not math.isfinite(a):
            raise ValueError("NONFINITE_ACCELERATION_NORM")
        by_side[side].append(a)
    spreads = {}
    for side in ("left", "right"):
        values = np.asarray(by_side[side], dtype=float)
        spreads[side] = float(np.percentile(values, 90) - np.percentile(values, 10))
        if spreads[side] < 0 or not math.isfinite(spreads[side]):
            raise ValueError("INVALID_ACCELERATION_DISPERSION")
    return abs(math.log((spreads["left"] + 1e-6) / (spreads["right"] + 1e-6)))


def main():
    target = {}
    for item in csv.DictReader((P36 / "TASK_INTERSECTION.csv").open(newline="")):
        if item["group"] == "PD":
            match = re.search(r"(\d{4}\.\d{2}\.\d{2})-\d{2}\.\d{2}\.\d{2}\.csv$", item["member_path"])
            assert match
            target[item["participant_id"]] = {"date": datetime.strptime(match.group(1), "%Y.%m.%d").date(), "path": item["member_path"]}
    valid_ids = {item["participant_id"] for item in csv.DictReader((P43 / "PERSON_STATUS.csv").open(newline="")) if item["status"] == "VALID"}
    assert len(target) == 44 and len(valid_ids) == 40 and valid_ids <= set(target)
    tail_start = reader.ZIP_SIZE - 4 * 1024 * 1024
    tail = reader.get_range(tail_start, reader.ZIP_SIZE - 1, 4 * 1024 * 1024)
    with zipfile.ZipFile(io.BytesIO(tail)) as archive:
        infos = {entry.filename: entry for entry in archive.infolist()}
    workbook_path = date_source.previous.WORKBOOK
    workbook_info = infos[workbook_path]
    raw = reader.get_csv_member(workbook_info, tail_start)
    sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)["Parkinson Disease"]
    source_rows = sheet.iter_rows(values_only=True)
    headers = list(next(source_rows))
    assert str(headers[13]).strip().lower() == "number of falls in the last month"
    assert str(headers[33]).strip().lower() == "average speed 10 m slow"
    entries = defaultdict(list)
    for row in source_rows:
        ident = date_source.previous.prior.canonical_id(row[0])
        if ident in valid_ids:
            entries[ident].append(row)
    selected = []
    source_reasons = Counter()
    for ident in sorted(valid_ids):
        rows = entries[ident]
        if not rows:
            source_reasons["ID_UNMATCHED"] += 1
            continue
        pairs = []
        for row in rows:
            dates = frozenset(date_source.dates_for_cell(row[1]))
            falls = date_source.previous.prior.numeric(row[13])
            speed = date_source.previous.prior.numeric(row[33])
            pairs.append((dates, falls, speed))
        if len(set(pairs)) > 1:
            source_reasons["DUPLICATE_CONFLICT"] += 1
            continue
        dates, falls, speed = pairs[0]
        close = {day for day in dates if abs((day - target[ident]["date"]).days) <= 7}
        if len(close) != 1 or falls is None or falls < 0 or int(falls) != falls:
            source_reasons["CROSSWALK_CHANGED_OR_INVALID"] += 1
            continue
        if speed is None or speed <= 0:
            source_reasons["SPEED_MISSING_OR_NONPOSITIVE"] += 1
            continue
        selected.append({"participant_id": ident, "y_fell_last_month": int(falls > 0), "log_clinical_speed": math.log(speed), "original_member_path": target[ident]["path"]})
    class_counts = dict(Counter(item["y_fell_last_month"] for item in selected))
    summary = {"run_id": "EXP044-DATA-001", "source_requests": 2, "requested_range_bytes": len(tail) + workbook_info.compress_size + 1024, "exp043_valid_people": 40, "speed_complete_people": len(selected), "speed_complete_class_counts": class_counts, "source_reasons": dict(source_reasons)}
    if len(selected) < 35 or class_counts.get(1, 0) < 8 or class_counts.get(0, 0) < 20:
        summary["gate"] = "STOP_MODEL_SOURCE"
        (HERE / "SOURCE_GATE.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
        print(json.dumps(summary, ensure_ascii=False))
        return
    summary["preflight"] = "PASS"
    failures = []
    for i, item in enumerate(selected, 1):
        info = infos[item["original_member_path"]]
        if info.compress_size > 1024 * 1024 - 1024 or info.file_size > 3 * 1024 * 1024:
            failures.append({"participant_id": item["participant_id"], "reason": "MEMBER_OUTSIDE_REGISTERED_BUDGET"})
            continue
        summary["source_requests"] += 1
        summary["requested_range_bytes"] += info.compress_size + 1024
        try:
            original = reader.get_csv_member(info, tail_start)
            item["bilateral_accel_asym"] = imu_asymmetry(original)
        except Exception as exc:
            failures.append({"participant_id": item["participant_id"], "reason": type(exc).__name__ + ": " + str(exc)[:200]})
        if i % 10 == 0 or i == len(selected):
            print(f"feature source {i}/{len(selected)}", flush=True)
    summary["feature_failures"] = failures
    if failures:
        summary["gate"] = "STOP_FEATURE_SOURCE"
    else:
        summary["gate"] = "ANALYSIS_SOURCE_READY"
        with (HERE / "ANALYSIS_TABLE.csv").open("x", newline="") as out:
            writer = csv.DictWriter(out, fieldnames=["participant_id", "y_fell_last_month", "log_clinical_speed", "bilateral_accel_asym"])
            writer.writeheader()
            writer.writerows({key: item[key] for key in writer.fieldnames} for item in selected)
    (HERE / "SOURCE_GATE.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(summary, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
