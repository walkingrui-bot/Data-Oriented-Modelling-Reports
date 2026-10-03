"""EXP042 explicit date-format candidates anchored to original recording dates."""

import csv
import io
import json
import math
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
P36 = ROOT / "STAT-PSYMOE-EXP036-20261002-001"
P37 = ROOT / "STAT-PSYMOE-EXP037-20261002-001"
P41 = ROOT / "STAT-PSYMOE-EXP041-20261002-001"
sys.path.insert(0, str(P37))
import audit_two_members as reader  # noqa: E402
sys.path.insert(0, str(P41))
import audit_clinical_crosswalk as prior  # noqa: E402

WORKBOOK = "Parkinson Dataset/Questionary data/Participants Data.xlsx"
FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d/%m/%Y", "%m/%d/%Y", "%d/%m/%y", "%m/%d/%y", "%d.%m.%Y", "%m.%d.%Y", "%d.%m.%y", "%m.%d.%y", "%d-%m-%Y", "%m-%d-%Y", "%d-%m-%y", "%m-%d-%y")


def format_pattern(value):
    if not isinstance(value, str):
        return type(value).__name__
    cleaned = value.strip()
    return re.sub(r"\d+", lambda m: "D" * len(m.group()), cleaned)[:50]


def candidate_days(value):
    if isinstance(value, datetime):
        return {value.date()}
    if isinstance(value, date):
        return {value}
    if isinstance(value, (int, float)) and math.isfinite(value) and 30000 <= value <= 60000:
        try:
            return {from_excel(value).date()}
        except (ValueError, OverflowError):
            return set()
    if not isinstance(value, str):
        return set()
    cleaned = value.strip()
    cleaned = re.sub(r"(?:[ T]\d{1,2}[:.]\d{2}(?::\d{2})?(?:\.\d+)?Z?)$", "", cleaned)
    output = set()
    for fmt in FORMATS:
        try:
            stamp = datetime.strptime(cleaned, fmt).date()
            if 2000 <= stamp.year <= 2099:
                output.add(stamp)
        except ValueError:
            pass
    return output


def main():
    targets = {}
    for row in csv.DictReader((P36 / "TASK_INTERSECTION.csv").open(newline="")):
        if row["group"] != "PD":
            continue
        match = re.search(r"(\d{4}\.\d{2}\.\d{2})-\d{2}\.\d{2}\.\d{2}\.csv$", row["member_path"])
        if match is None:
            raise RuntimeError("original PD recording date missing")
        targets[row["participant_id"]] = datetime.strptime(match.group(1), "%Y.%m.%d").date()
    assert len(targets) == 44
    tail_start = reader.ZIP_SIZE - 4 * 1024 * 1024
    tail = reader.get_range(tail_start, reader.ZIP_SIZE - 1, 4 * 1024 * 1024)
    with zipfile.ZipFile(io.BytesIO(tail)) as archive:
        info = {entry.filename: entry for entry in archive.infolist()}[WORKBOOK]
    raw = reader.get_csv_member(info, tail_start)
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = book["Parkinson Disease"]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    falls_columns = [i for i, v in enumerate(headers) if str(v or "").strip().lower() == "number of falls in the last month"]
    assert falls_columns == [13]
    by_id = defaultdict(list)
    for line, cells in enumerate(rows, 2):
        ident = prior.canonical_id(cells[0])
        if ident in targets:
            by_id[ident].append((line, cells[1], cells[13]))
    patterns = Counter()
    statuses = []
    values = []
    for ident in sorted(targets):
        entries = by_id[ident]
        if not entries:
            status = "ID_UNMATCHED"
        else:
            for _, date_cell, _ in entries:
                patterns[format_pattern(date_cell)] += 1
            pairs = [(frozenset(candidate_days(date_cell)), prior.numeric(fall_cell)) for _, date_cell, fall_cell in entries]
            if len(set(pairs)) > 1:
                status = "DUPLICATE_CONFLICT"
            else:
                candidate_set, outcome = pairs[0]
                close_days = {day for day in candidate_set if abs((day - targets[ident]).days) <= 7}
                if len(close_days) == 0:
                    status = "DATE_NO_UNIQUE_MATCH_IN_7_DAYS"
                elif len(close_days) > 1:
                    status = "DATE_AMBIGUOUS_NEAR_RECORDING"
                elif outcome is None or outcome < 0 or int(outcome) != outcome:
                    status = "FALLS_MISSING_OR_INVALID"
                else:
                    status = "VALID"
                    values.append(int(outcome))
        statuses.append({"participant_id": ident, "workbook_row_count": len(entries), "status": status})
    status_counts = dict(Counter(r["status"] for r in statuses))
    outcome_hist = dict(sorted(Counter(values).items()))
    audit = {"run_id": "EXP042-DATE-001", "source_member": WORKBOOK, "source_requests": 2, "requested_range_bytes": len(tail) + info.compress_size + 1024, "fixed_PD_people": 44, "date_format_patterns": dict(patterns), "status_counts": status_counts, "valid_distinct_people": len(values), "distinct_fall_counts": len(set(values)), "aggregate_fall_histogram": outcome_hist, "gate": "FALLS_OUTCOME_SOURCE_READY" if len(values) >= 30 and len(set(values)) >= 5 else "STOP_FALLS_OUTCOME_SUPPORT"}
    (HERE / "DATE_FORMAT_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n")
    with (HERE / "PERSON_STATUS.csv").open("x", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["participant_id", "workbook_row_count", "status"])
        writer.writeheader()
        writer.writerows(statuses)
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
