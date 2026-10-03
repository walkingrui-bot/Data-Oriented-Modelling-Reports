"""EXP043 explicit English month date audit on the same fixed clinical outcome."""

import csv
import io
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
P36 = ROOT / "STAT-PSYMOE-EXP036-20261002-001"
P42 = ROOT / "STAT-PSYMOE-EXP042-20261002-001"
sys.path.insert(0, str(P42))
import audit_date_formats as previous  # noqa: E402

MONTHS = {name.lower(): i for i, name in enumerate(("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"), 1)}


def dates_for_cell(value):
    days = previous.candidate_days(value)
    if isinstance(value, str):
        match = re.fullmatch(r"([A-Za-z]+)\s+(\d{1,2}),\s*(\d{4})", value.strip())
        if match and match.group(1).lower() in MONTHS:
            try:
                days.add(date(int(match.group(3)), MONTHS[match.group(1).lower()], int(match.group(2))))
            except ValueError:
                pass
    return days


def main():
    targets = {}
    for item in csv.DictReader((P36 / "TASK_INTERSECTION.csv").open(newline="")):
        if item["group"] == "PD":
            match = re.search(r"(\d{4}\.\d{2}\.\d{2})-\d{2}\.\d{2}\.\d{2}\.csv$", item["member_path"])
            assert match
            targets[item["participant_id"]] = datetime.strptime(match.group(1), "%Y.%m.%d").date()
    assert len(targets) == 44
    reader = previous.reader
    tail_start = reader.ZIP_SIZE - 4 * 1024 * 1024
    tail = reader.get_range(tail_start, reader.ZIP_SIZE - 1, 4 * 1024 * 1024)
    with zipfile.ZipFile(io.BytesIO(tail)) as archive:
        info = {entry.filename: entry for entry in archive.infolist()}[previous.WORKBOOK]
    raw = reader.get_csv_member(info, tail_start)
    sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)["Parkinson Disease"]
    source = sheet.iter_rows(values_only=True)
    header = next(source)
    assert str(header[13]).strip().lower() == "number of falls in the last month"
    by_id = defaultdict(list)
    for line, row in enumerate(source, 2):
        ident = previous.prior.canonical_id(row[0])
        if ident in targets:
            by_id[ident].append((line, row[1], row[13]))
    values = []
    statuses = []
    english_month_total = 0
    english_month_valid = 0
    for ident in sorted(targets):
        entries = by_id[ident]
        for _, date_cell, _ in entries:
            if isinstance(date_cell, str) and re.fullmatch(r"[A-Za-z]+\s+\d{1,2},\s*\d{4}", date_cell.strip()):
                english_month_total += 1
        if not entries:
            status = "ID_UNMATCHED"
        else:
            pairs = [(frozenset(dates_for_cell(date_cell)), previous.prior.numeric(falls)) for _, date_cell, falls in entries]
            if len(set(pairs)) > 1:
                status = "DUPLICATE_CONFLICT"
            else:
                day_set, falls = pairs[0]
                close = {day for day in day_set if abs((day - targets[ident]).days) <= 7}
                if len(close) == 0:
                    status = "DATE_NO_UNIQUE_MATCH_IN_7_DAYS"
                elif len(close) > 1:
                    status = "DATE_AMBIGUOUS_NEAR_RECORDING"
                elif falls is None or falls < 0 or int(falls) != falls:
                    status = "FALLS_MISSING_OR_INVALID"
                else:
                    status = "VALID"
                    values.append(int(falls))
                    if isinstance(entries[0][1], str) and re.fullmatch(r"[A-Za-z]+\s+\d{1,2},\s*\d{4}", entries[0][1].strip()):
                        english_month_valid += 1
        statuses.append({"participant_id": ident, "workbook_row_count": len(entries), "status": status})
    counts = dict(Counter(r["status"] for r in statuses))
    hist = dict(sorted(Counter(values).items()))
    result = {"run_id": "EXP043-MONTH-001", "source_member": previous.WORKBOOK, "source_requests": 2, "requested_range_bytes": len(tail) + info.compress_size + 1024, "fixed_PD_people": 44, "english_month_rows": english_month_total, "english_month_valid_outcomes": english_month_valid, "status_counts": counts, "valid_distinct_people": len(values), "distinct_fall_counts": len(set(values)), "aggregate_fall_histogram": hist, "gate": "FALLS_OUTCOME_SOURCE_READY_FOR_DESIGN" if len(values) >= 30 and len(set(values)) >= 5 else "STOP_FALLS_OUTCOME_SUPPORT"}
    (HERE / "MONTH_DATE_AUDIT.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    with (HERE / "PERSON_STATUS.csv").open("x", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["participant_id", "workbook_row_count", "status"])
        writer.writeheader()
        writer.writerows(statuses)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
