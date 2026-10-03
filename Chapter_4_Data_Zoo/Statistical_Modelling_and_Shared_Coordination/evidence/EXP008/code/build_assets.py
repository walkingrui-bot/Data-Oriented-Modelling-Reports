import os, json, hashlib, shutil, zipfile
import numpy as np, pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT='/mnt/data/INTERNAL_COORDINATION_008'; RES=ROOT+'/results'; FIG=ROOT+'/figures'; os.makedirs(FIG,exist_ok=True)
df=pd.read_csv(RES+'/per_seed.csv')
def sub(mode='aligned',dyn=1): return df[(df['mode']==mode)&(df.dynamic==dyn)]
def med(mode,dyn,col): return float(sub(mode,dyn)[col].median())
def pct(mode,dyn,num,den):
 d=sub(mode,dyn); return float(np.median((d[num]/d[den]-1)*100))
# figures
plt.figure(figsize=(7.2,4.6))
labels=['Aligned\ncorrect time','Aligned\nreversed time','Event baseline\ncorrect','Event baseline\nreversed']
vals=[med('aligned',1,'correct_pinball'),med('aligned',1,'reverse_time_pinball'),med('event',1,'correct_pinball'),med('event',1,'reverse_time_pinball')]
plt.bar(labels,vals); plt.ylabel('Pinball loss (lower is better)'); plt.title('Dynamic current-state evidence: chronology matters'); plt.tight_layout(); plt.savefig(FIG+'/fig1_chronology.png',dpi=180); plt.close()
plt.figure(figsize=(7.2,4.6))
labels=['Aligned set update','Serial event update']; vals=[med('aligned',1,'within_slice_pred_sd'),med('event',1,'within_slice_pred_sd')]
plt.bar(labels,vals); plt.ylabel('Prediction SD across within-slice permutations'); plt.title('Same-time arrival order'); plt.tight_layout(); plt.savefig(FIG+'/fig2_same_time_order.png',dpi=180); plt.close()
plt.figure(figsize=(7.2,4.6))
labels=['Dynamic correct','Dynamic half stale','Static correct','Static half stale']
vals=[med('aligned',1,'correct_pinball'),med('aligned',1,'stale_half_pinball'),med('aligned',0,'correct_pinball'),med('aligned',0,'stale_half_pinball')]
plt.bar(labels,vals); plt.ylabel('Pinball loss'); plt.title('Timestamp alignment determines which evidence describes the current slice'); plt.tight_layout(); plt.savefig(FIG+'/fig3_misalignment.png',dpi=180); plt.close()
plt.figure(figsize=(7.2,4.6))
labels=['Aligned ganglion','Serial event baseline']; vals=[med('aligned',1,'latent_R2'),med('event',1,'latent_R2')]
plt.bar(labels,vals); plt.ylim(.9,1.0); plt.ylabel('Held-out linear R² to current world state'); plt.title('Audit-only current-state readout'); plt.tight_layout(); plt.savefig(FIG+'/fig4_state_readout.png',dpi=180); plt.close()

def setup(doc):
 sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
 for name,size in [('Normal',10.2),('Title',25),('Heading 1',16),('Heading 2',12.5)]:
  st=doc.styles[name]; st.font.name='Aptos Display' if name!='Normal' else 'Aptos'; st.font.size=Pt(size)
 if 'Caption Small' not in [s.name for s in doc.styles]:
  st=doc.styles.add_style('Caption Small',1); st.font.name='Aptos'; st.font.size=Pt(8.5); st.font.italic=True
 return doc
def shade(cell,fill='EDEDED'):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def table(doc,headers,rows):
 t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i]; c.text=str(h); shade(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for r in c.paragraphs[0].runs: r.font.bold=True; r.font.size=Pt(8.3)
 for row in rows:
  cs=t.add_row().cells
  for i,v in enumerate(row):
   cs[i].text=str(v); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for r in cs[i].paragraphs[0].runs: r.font.size=Pt(8.3)
 return t
def fig(doc,path,cap,w=6.2):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(path,width=Inches(w)); cp=doc.add_paragraph(style='Caption Small'); cp.add_run(cap)

D=setup(Document()); D.add_paragraph('INTERNAL COORDINATION 008',style='Title'); D.add_paragraph('Time-Aligned Evidence Coordination',style='Subtitle'); D.add_paragraph('Same-time commutativity and across-time state separation | 1 October 2026')
D.add_heading('Result in one paragraph',1)
D.add_paragraph(f"Experiment 008 gives the recurrent shared ganglion an explicit temporal hierarchy without asking it to predict the future. Evidence carrying the same timestamp is first coordinated as an unordered set and then passed once through the shared 16×16 ganglion. Evidence from later timestamps updates the ganglion recurrently as a new current-reality slice. Across five seeds, permuting six same-time channels changed predicted medians only at numerical precision (median SD {med('aligned',1,'within_slice_pred_sd'):.2e}). In dynamic episodes, reversing the three real time slices increased current-state pinball loss from {med('aligned',1,'correct_pinball'):.4f} to {med('aligned',1,'reverse_time_pinball'):.4f}; in static episodes the same reversal left the score essentially unchanged ({med('aligned',0,'correct_pinball'):.4f} versus {med('aligned',0,'reverse_time_pinball'):.4f}). Mislabeling half of the oldest measurements as current increased dynamic loss to {med('aligned',1,'stale_half_pinball'):.4f}. The audit-only linear readout from the final shared state recovered the actual current 10-dimensional world state with median held-out R² {med('aligned',1,'latent_R2'):.3f}.")
D.add_heading('1. What changed from Experiment 007',1)
D.add_paragraph('Experiment 007 treated every current observation as a serial recurrent event, which made the final state depend on arbitrary presentation order even when all observations described the same time slice. Experiment 008 distinguishes two kinds of ordering. Within one timestamp, channel observations form a set. Across timestamps, the shared ganglion remains recurrent, because the evidence refers to different current realities.')
table(D,['Level','Update rule','Meaning'],[
 ['Same timestamp','Encode each available channel, average the 16-D evidence vectors, apply the shared ganglion once','Arrival order carries no reality information'],
 ['Later timestamp','Apply the same shared ganglion again to the new time-slice evidence','Chronology carries reality information when the world changed'],
 ['Output target','Five conditional quantiles of an independent repeat measurement at the current timestamp','Current-state distribution; no future target'],
])
D.add_heading('2. Data construction and training',1)
D.add_paragraph('The experiment keeps the real 10-variable patient geometry used in Experiments 005–007 and constructs controlled three-slice current-world trajectories along principal directions estimated from the training patients. Each time slice is then observed through the same six epidemiology-inspired stochastic measurement channels used previously. Twenty-five percent of training episodes are static; the remainder change across the three slices. Independent source and target measurements are drawn at every slice. The model is scored only against the current slice after evidence from that slice arrives.')
D.add_paragraph('The aligned architecture uses the channel input neurons to produce 16-dimensional evidence vectors. All evidence carrying the same time label is pooled commutatively. The resulting slice vector and the previous reality state pass through the same trainable 16×16 ganglion. Channel-specific output neurons return five ordered quantiles for current repeat measurements.')
D.add_heading('3. Same-time evidence is commutative',1); fig(D,FIG+'/fig2_same_time_order.png','Figure 1. Prediction variation after permuting the six channels inside every time slice. Set coordination removes arbitrary same-time presentation order; the serial-event baseline retains a large order effect.')
table(D,['Architecture','Dynamic current pinball','Same-time prediction SD','Same-time state SD'],[
 ['Time-aligned set update',f"{med('aligned',1,'correct_pinball'):.4f}",f"{med('aligned',1,'within_slice_pred_sd'):.2e}",f"{med('aligned',1,'within_slice_state_sd'):.2e}"],
 ['Serial event baseline',f"{med('event',1,'correct_pinball'):.4f}",f"{med('event',1,'within_slice_pred_sd'):.4f}",f"{med('event',1,'within_slice_state_sd'):.4f}"],
])
D.add_paragraph('The near-zero aligned values are architectural: evidence with the same time label is combined before the recurrent ganglion update. This directly encodes the statement that arbitrary arrival order within one reality slice should not create a new reality.')
D.add_heading('4. Across-time order is intentionally non-commutative',1); fig(D,FIG+'/fig1_chronology.png','Figure 2. Dynamic episodes scored against the true current time slice. Reversing real time slices changes the inferred current reality and worsens the proper distributional score.')
table(D,['Episode','Correct chronology','Reversed chronology','Median penalty'],[
 ['Dynamic world',f"{med('aligned',1,'correct_pinball'):.4f}",f"{med('aligned',1,'reverse_time_pinball'):.4f}",f"{pct('aligned',1,'reverse_time_pinball','correct_pinball'):.1f}%"],
 ['Static world',f"{med('aligned',0,'correct_pinball'):.4f}",f"{med('aligned',0,'reverse_time_pinball'):.4f}",f"{pct('aligned',0,'reverse_time_pinball','correct_pinball'):.2f}%"],
])
D.add_paragraph(f"For dynamic episodes, the final shared state moves substantially when the actual time slices are reversed (median state distance {med('aligned',1,'reverse_state_distance'):.3f}). For static episodes, the distributional score is invariant to the same reversal even though independent measurement noise gives each slice slightly different observations. The experiment therefore separates chronology from arbitrary presentation order: time order matters to the extent that reality itself changed.")
D.add_heading('5. Misaligned timestamps create a different current reality',1); fig(D,FIG+'/fig3_misalignment.png','Figure 3. Half of the oldest measurements are relabeled as current. Dynamic episodes are displaced; static episodes remain nearly unchanged.')
table(D,['Condition','Correct current slice','Half of current channels actually stale','State displacement'],[
 ['Dynamic',f"{med('aligned',1,'correct_pinball'):.4f}",f"{med('aligned',1,'stale_half_pinball'):.4f}",f"{med('aligned',1,'stale_half_state_distance'):.3f}"],
 ['Static',f"{med('aligned',0,'correct_pinball'):.4f}",f"{med('aligned',0,'stale_half_pinball'):.4f}",f"{med('aligned',0,'stale_half_state_distance'):.3f}"],
])
D.add_paragraph('This test makes timestamp alignment operational. Measurements from different realities should not be pooled merely because they share a channel type or arrive together in memory. When reality has moved, stale evidence labeled as current changes the inferred current state. When reality is unchanged, the same relabeling carries little cost.')
D.add_heading('6. Current-state audit',1); fig(D,FIG+'/fig4_state_readout.png','Figure 4. Audit-only linear recovery of the actual current patient-state vector from the final internal state. Simulator/current-world values never enter the training objective.')
D.add_paragraph(f"The time-aligned shared state gives median held-out R² {med('aligned',1,'latent_R2'):.3f} for the actual current 10-dimensional world state, compared with {med('event',1,'latent_R2'):.3f} for the serial-event baseline. The model was trained only on stochastic current measurement distributions; the world-state coordinates are used after training for audit.")
D.add_heading('7. Engineering interpretation',1)
D.add_paragraph('The resulting circuit has two coordination scales. Channel input neurons convert native measurements into evidence vectors. Evidence with one timestamp is coordinated as a set. The shared ganglion then carries the resulting current reality across genuine time slices. Output neurons express that current reality as channel-specific measurement distributions. The same shared matrix therefore participates in both within-slice evidence integration and across-slice reality updating, while the timestamp boundary decides which operation is appropriate.')
D.add_heading('Next experiment interface',1)
D.add_paragraph('The next experiment can remove the clean global slice boundary and admit asynchronous channel timestamps. The coordination system would then receive explicit observation times, combine only evidence whose temporal support overlaps, and maintain separate current-reality updates when evidence refers to distinct times. Evaluation should continue to score observed current-time distributions rather than future predictions.')
D.add_heading('Version history',1); D.add_paragraph('Version 0.1 establishes five-seed time-aligned evidence coordination, same-time permutation invariance by construction, dynamic versus static chronology controls, timestamp-misalignment tests, current-distribution calibration, and an audit-only current-world readout.')
report=ROOT+'/INTERNAL_COORDINATION_008_Report_v0.1_20261001.docx'; D.save(report)
# append to living 0.7
src='/mnt/data/INTERNAL_COORDINATION_007/INTERNAL_COORDINATION_Living_Report_v0.7_20261001.docx'; L=Document(src)
# update visible version marker when present
for p in L.paragraphs[:12]:
 if p.text.startswith('Version 0.7'): p.text='Version 0.8   1 October 2026'
sec=L.add_section(WD_SECTION.NEW_PAGE); sec.footer.is_linked_to_previous=False; sec.footer.paragraphs[0].text='Experiment 008'
L.add_heading('Experiment 008 | Time-Aligned Evidence Coordination',1); L.add_paragraph('Same-time commutativity and across-time current-reality separation')
L.add_paragraph(f"Evidence carrying one timestamp is now pooled before the shared 16×16 recurrent ganglion update. Evidence from later timestamps updates the ganglion again as a new current-reality slice. Across five seeds, same-time channel permutations changed predicted medians only at numerical precision ({med('aligned',1,'within_slice_pred_sd'):.2e}). In dynamic episodes, reversing the true three-slice chronology increased current pinball loss from {med('aligned',1,'correct_pinball'):.4f} to {med('aligned',1,'reverse_time_pinball'):.4f}; in static episodes the same reversal left the score unchanged ({med('aligned',0,'correct_pinball'):.4f} versus {med('aligned',0,'reverse_time_pinball'):.4f}).")
fig(L,FIG+'/fig1_chronology.png','Figure 008-1. Real chronology is retained only when the underlying world changes.',w=6.05)
L.add_heading('Timestamp alignment',2); fig(L,FIG+'/fig3_misalignment.png','Figure 008-2. Stale evidence relabeled as current changes the inferred reality only when the world has actually moved.',w=6.05)
L.add_paragraph(f"Relabeling half of the oldest measurements as current raises dynamic loss to {med('aligned',1,'stale_half_pinball'):.4f} and displaces the final state by {med('aligned',1,'stale_half_state_distance'):.3f}; the same operation in static episodes leaves the score essentially unchanged. The audit-only current-state readout reaches held-out R² {med('aligned',1,'latent_R2'):.3f}.")
L.add_heading('Engineering consequence',2); L.add_paragraph('Order should be removed only inside a shared temporal support. Chronology remains part of the model when observations refer to different realities. The shared ganglion therefore receives unordered evidence sets at each aligned time slice and recurrently carries the current reality between slices. Training and evaluation remain current-state distributional; no future observation is generated as an internal training target.')
L.add_paragraph('Version history: v0.8 adds Experiment 008 with explicit temporal support, same-time set coordination, dynamic/static chronology controls, timestamp-misalignment tests, and current-world audit.')
live=ROOT+'/INTERNAL_COORDINATION_Living_Report_v0.8_20261001.docx'; L.save(live)
# protocol / readme / manifest / zip
protocol={'experiment':'INTERNAL_COORDINATION_008','date':'2026-10-01','seeds':list(range(5)),'world_slices':3,'channels':6,'variables_per_channel':10,'ganglion_dim':16,'quantiles':[0.1,0.25,0.5,0.75,0.9],'train_steps':420,'train_static_fraction':0.25,'training_target':'independent repeat measurement distribution at the current timestamp','same_time_operation':'commutative mean of channel evidence before one ganglion update','across_time_operation':'recurrent ganglion update per distinct time slice'}
json.dump(protocol,open(ROOT+'/protocol.json','w'),indent=2)
open(ROOT+'/README.md','w').write('''# INTERNAL_COORDINATION_008 — Time-Aligned Evidence Coordination\n\nThis experiment extends the recurrent shared-ganglion architecture with an explicit temporal hierarchy. Measurements sharing a timestamp are coordinated as an unordered evidence set. Distinct timestamps update the same ganglion recurrently. The training target is an independent repeat-measurement distribution at the current timestamp; no future prediction target is used.\n\nReproduce:\n```bash\npython code/time_aligned_coordination.py\npython code/build_assets.py\n```\n\nKey outputs: `results/per_seed.csv`, four figures, standalone report, and Living Report v0.8. `data/diabetes_backbone.npz` is the fixed real-patient scaffold inherited from the previous experiments.\n''')
# manifest
files=[]
for base,dirs,fs in os.walk(ROOT):
 for f in fs:
  if f.endswith('.zip') or '/render_' in base: continue
  path=os.path.join(base,f); rel=os.path.relpath(path,ROOT)
  with open(path,'rb') as h: dat=h.read()
  files.append({'path':rel,'bytes':len(dat),'sha256':hashlib.sha256(dat).hexdigest()})
json.dump({'files':files},open(ROOT+'/manifest.json','w'),indent=2)
zip_path='/mnt/data/INTERNAL_COORDINATION_008_Evidence.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
 for x in files: z.write(os.path.join(ROOT,x['path']),arcname='INTERNAL_COORDINATION_008/'+x['path'])
print(report); print(live); print(zip_path)
