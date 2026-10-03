"""EXP041 fixed candidate clinical outcome audit; no raw workbook copy."""

import csv
import io
import json
import math
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
P36 = ROOT / "STAT-PSYMOE-EXP036-20261002-001"
P37 = ROOT / "STAT-PSYMOE-EXP037-20261002-001"
sys.path.insert(0, str(P37))
import audit_two_members as reader  # noqa: E402

WORKBOOK = "Parkinson Dataset/Questionary data/Participants Data.xlsx"
DATE_RE = re.compile(r"(\d{4}\.\d{2}\.\d{2})-(\d{2}\.\d{2}\.\d{2})\.csv$")


def canonical_id(value):
    if isinstance(value, (int, float)) and math.isfinite(value) and int(value) == value:
        return f"{int(value):02d}"
    if isinstance(value, str):
        cleaned = value.strip().upper()
        if cleaned.isdigit():
            return f"{int(cleaned):02d}"
    return None


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and math.isfinite(value) and 30000 <= value <= 60000:
        try:
            return from_excel(value).date()
        except (ValueError, OverflowError):
            return None
    if isinstance(value, str):
        cleaned = value.strip()
        for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(cleaned, fmt).date()
            except ValueError:
                pass
    return None


def numeric(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    elif isinstance(value, str):
        try:
            number = float(value.strip().replace(",", "."))
        except ValueError:
            return None
    else:
        return None
    return number if math.isfinite(number) else None


def field_candidates(headers):
    detected = {"UPDRS_III": [], "MINI_BEST_TOTAL": [], "FALLS_MONTH": []}
    for index, original in enumerate(headers):
        title = " ".join(str(original or "").lower().replace("-", " ").split())
        if ("updrs" in title and ("iii" in title or "part 3" in title or "part iii" in title) and ("total" in title or "score" in title)):
            detected["UPDRS_III"].append((index, original))
        if "mini" in title and "best" in title and ("total" in title or "score" in title):
            detected["MINI_BEST_TOTAL"].append((index, original))
        if "falls" in title and "last month" in title and ("number" in title or "count" in title):
            detected["FALLS_MONTH"].append((index, original))
    return detected


def main():
    source_rows = list(csv.DictReader((P36 / "TASK_INTERSECTION.csv").open(newline="")))
    targets = {}
    for row in source_rows:
        if row["group"] != "PD":
            continue
        match = DATE_RE.search(row["member_path"])
        if not match:
            raise RuntimeError("PD file date absent")
        targets[row["participant_id"]] = datetime.strptime(match.group(1), "%Y.%m.%d").date()
    if len(targets) != 44:
        raise RuntimeError("fixed PD target does not have 44 IDs")
    tail_start = reader.ZIP_SIZE - 4 * 1024 * 1024
    tail = reader.get_range(tail_start, reader.ZIP_SIZE - 1, 4 * 1024 * 1024)
    with zipfile.ZipFile(io.BytesIO(tail)) as archive:
        info = {entry.filename: entry for entry in archive.infolist()}[WORKBOOK]
    raw = reader.get_csv_member(info, tail_start)
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = book["Parkinson Disease"]
    source = sheet.iter_rows(values_only=True)
    headers = list(next(source))
    candidates = field_candidates(headers)
    by_id = defaultdict(list)
    id_invalid_count = 0
    date_types = Counter()
    for row_number, row in enumerate(source, 2):
        if not any(value is not None for value in row):
            continue
        ident = canonical_id(row[0])
        if ident is None:
            id_invalid_count += 1
            continue
        if ident not in targets:
            continue
        date_types[type(row[1]).__name__] += 1
        by_id[ident].append({"row_number": row_number, "date": parse_date(row[1]), "values": row})
    outcome_audit = {}
    status_by_person = {ident: {"participant_id": ident, "workbook_row_count": len(by_id[ident])} for ident in sorted(targets)}
    for name in ("UPDRS_III", "MINI_BEST_TOTAL", "FALLS_MONTH"):
        fields = candidates[name]
        if len(fields) != 1:
            outcome_audit[name] = {"matching_columns": [{"index_1_based": i + 1, "name": str(label)} for i, label in fields], "column_status": "ABSENT" if not fields else "AMBIGUOUS_MULTIPLE", "valid_distinct_people": 0, "distinct_numeric_values": 0}
            for ident in status_by_person:
                status_by_person[ident][name] = "FIELD_ABSENT_OR_AMBIGUOUS"
            continue
        col, label = fields[0]
        valid_values = []
        statuses = Counter()
        for ident, record in status_by_person.items():
            entries = by_id[ident]
            if not entries:
                status = "ID_UNMATCHED"
            else:
                pairs = [(entry["date"], numeric(entry["values"][col] if col < len(entry["values"]) else None)) for entry in entries]
                if len(set(pairs)) > 1:
                    status = "DUPLICATE_CONFLICT"
                elif pairs[0][0] is None:
                    status = "DATE_INVALID"
                elif abs((pairs[0][0] - targets[ident]).days) > 7:
                    status = "DATE_OUTSIDE_7_DAYS"
                elif pairs[0][1] is None:
                    status = "OUTCOME_MISSING_OR_NONNUMERIC"
                else:
                    status = "VALID"
                    valid_values.append(pairs[0][1])
            record[name] = status
            statuses[status] += 1
        outcome_audit[name] = {"matching_columns": [{"index_1_based": col + 1, "name": str(label)}], "column_status": "UNIQUE", "status_counts": dict(statuses), "valid_distinct_people": len(valid_values), "distinct_numeric_values": len(set(valid_values)), "value_min": min(valid_values) if valid_values else None, "value_max": max(valid_values) if valid_values else None, "passes_support": len(valid_values) >= 30 and len(set(valid_values)) >= 5}
    selected = next((name for name in ("UPDRS_III", "MINI_BEST_TOTAL", "FALLS_MONTH") if outcome_audit[name].get("passes_support")), None)
    audit = {"run_id": "EXP041-CROSSWALK-001", "source_member": WORKBOOK, "source_requests": 2, "requested_range_bytes": len(tail) + info.compress_size + 1024, "fixed_PD_people": 44, "workbook_sheet": "Parkinson Disease", "workbook_data_rows": sheet.max_row - 1, "unparseable_id_rows": id_invalid_count, "matched_target_IDs_with_any_row": sum(bool(by_id[i]) for i in targets), "date_cell_types_for_matched_rows": dict(date_types), "candidate_outcomes": outcome_audit, "selected_outcome": selected, "gate": "PD_CLINICAL_CROSSWALK_READY" if selected else "STOP_CLINICAL_OUTCOME_SUPPORT"}
    (HERE / "CLINICAL_FIELD_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n")
    with (HERE / "PERSON_MATCH_STATUS.csv").open("x", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["participant_id", "workbook_row_count", "UPDRS_III", "MINI_BEST_TOTAL", "FALLS_MONTH"])
        writer.writeheader()
        writer.writerows(status_by_person.values())
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
