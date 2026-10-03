"""EXP040 date and companion metadata audit with source ZIP members in memory."""

import csv
import io
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
P36 = ROOT / "STAT-PSYMOE-EXP036-20261002-001"
P37 = ROOT / "STAT-PSYMOE-EXP037-20261002-001"
sys.path.insert(0, str(P37))
import audit_two_members as reader  # noqa: E402

README = "Parkinson Dataset/README.docx"
WORKBOOK = "Parkinson Dataset/Questionary data/Participants Data.xlsx"
DATE_RE = re.compile(r"(\d{4}\.\d{2}\.\d{2})-(\d{2}\.\d{2}\.\d{2})\.csv$")


def audit_paths():
    target = list(csv.DictReader((P36 / "TASK_INTERSECTION.csv").open(newline="")))
    assert len(target) == 89 and Counter(r["group"] for r in target) == {"PD": 44, "HC": 45}
    dates = defaultdict(list)
    failures = []
    for item in target:
        match = DATE_RE.search(item["member_path"])
        if not match:
            failures.append({"group": item["group"], "participant_id": item["participant_id"], "reason": "FILENAME_DATE_ABSENT"})
            continue
        try:
            stamp = datetime.strptime(match.group(1) + "-" + match.group(2), "%Y.%m.%d-%H.%M.%S")
        except ValueError:
            failures.append({"group": item["group"], "participant_id": item["participant_id"], "reason": "FILENAME_DATE_INVALID"})
            continue
        dates[item["group"]].append(stamp)
    months = {group: dict(sorted(Counter(d.strftime("%Y-%m") for d in dates[group]).items())) for group in ("PD", "HC")}
    all_months = sorted(set(months["PD"]) | set(months["HC"]))
    by_month = [{"month": m, "PD": months["PD"].get(m, 0), "HC": months["HC"].get(m, 0)} for m in all_months]
    sufficient = [row for row in by_month if row["PD"] >= 10 and row["HC"] >= 10]
    return {"source_path_table": str(P36 / "TASK_INTERSECTION.csv"), "filename_date_parse_failures": failures, "date_semantics": "FILENAME_TOKEN_ONLY_PENDING_README", "by_group": {g: {"count": len(dates[g]), "earliest_filename_timestamp": min(dates[g]).isoformat() if dates[g] else None, "latest_filename_timestamp": max(dates[g]).isoformat() if dates[g] else None, "by_month": months[g]} for g in ("PD", "HC")}, "month_cross_tab": by_month, "months_with_at_least_10_each": sufficient}


def main():
    dates = audit_paths()
    tail_start = reader.ZIP_SIZE - 4 * 1024 * 1024
    tail = reader.get_range(tail_start, reader.ZIP_SIZE - 1, 4 * 1024 * 1024)
    with zipfile.ZipFile(io.BytesIO(tail)) as archive:
        infos = {info.filename: info for info in archive.infolist()}
    raw_readme = reader.get_csv_member(infos[README], tail_start)
    raw_workbook = reader.get_csv_member(infos[WORKBOOK], tail_start)
    document = Document(io.BytesIO(raw_readme))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    book = load_workbook(io.BytesIO(raw_workbook), read_only=True, data_only=True)
    sheets = []
    for sheet in book.worksheets:
        candidates = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 4), values_only=True):
            cells = list(row[:50])
            candidates.append({"text_cells": sum(isinstance(v, str) for v in cells), "numeric_cells": sum(isinstance(v, (int, float)) and not isinstance(v, bool) for v in cells), "empty_cells": sum(v is None for v in cells)})
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        first_labels = [str(v)[:80] if isinstance(v, str) else None for v in first[:50]]
        sheets.append({"name": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "first_row_field_labels": first_labels, "first_four_row_type_counts": candidates})
    metadata = {"readme_original_member": README, "workbook_original_member": WORKBOOK, "readme_paragraph_count": len(paragraphs), "readme_table_count": len(document.tables), "workbook_sheets": sheets, "source_requests": 3, "requested_range_bytes": len(tail) + infos[README].compress_size + infos[WORKBOOK].compress_size + 2048}
    (HERE / "DATE_SUPPORT.json").write_text(json.dumps(dates, ensure_ascii=False, indent=2) + "\n")
    (HERE / "METADATA_SCHEMA.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n")
    print("DATES", json.dumps({"by_group": dates["by_group"], "months_with_at_least_10_each": dates["months_with_at_least_10_each"]}, ensure_ascii=False))
    print("README_PARAGRAPHS", len(paragraphs))
    for i, paragraph in enumerate(paragraphs):
        print("README", i, paragraph[:500])
    print("SHEETS", [(s["name"], s["max_row"], s["max_column"]) for s in sheets])


if __name__ == "__main__":
    main()
