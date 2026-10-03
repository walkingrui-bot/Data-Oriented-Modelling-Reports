"""Inspect official FDA files in memory; persist metadata, never source copies."""

from __future__ import annotations

import csv
import io
import json
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    "drugs_at_fda": "https://www.fda.gov/media/89850/download?attachment",
    "orange_book": "https://www.fda.gov/media/76860/download?attachment",
    "purple_book_august_2026": "https://www.accessdata.fda.gov/drugsatfda_docs/PurpleBook/2026/purplebook-search-August-data-download.csv",
    "purple_book_august_2026_xlsx": "https://www.accessdata.fda.gov/drugsatfda_docs/PurpleBook/2026/purplebook-search-August-data-download.xlsx",
}


def download(url: str) -> tuple[bytes, dict]:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 EXP021 research"})
    with urllib.request.urlopen(request, timeout=90) as response:
        blob = response.read()
        meta = {
            "source_url": url,
            "resolved_url": response.url,
            "http_status": response.status,
            "content_type": response.headers.get("Content-Type"),
            "last_modified": response.headers.get("Last-Modified"),
            "bytes_read": len(blob),
            "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
        }
    return blob, meta


def inspect_zip(blob: bytes, delimiter: str) -> dict:
    result = {}
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        for member in archive.namelist():
            with archive.open(member) as stream:
                rows = csv.reader(io.TextIOWrapper(stream, encoding="utf-8-sig", errors="replace", newline=""), delimiter=delimiter)
                header = next(rows, [])
                count = sum(1 for _ in rows)
            result[member] = {"columns": header, "rows_after_header": count}
    return result


def inspect_purple(blob: bytes) -> dict:
    rows = list(csv.reader(io.StringIO(blob.decode("utf-8-sig", errors="replace"))))
    markers = [(i, row[0]) for i, row in enumerate(rows) if row and row[0].startswith(("N/R/U", "Applicant"))]
    return {"physical_rows": len(rows), "section_headers": markers, "first_section_header": rows[3] if len(rows) > 3 else []}


def inspect_purple_xlsx(blob: bytes) -> dict:
    workbook = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    sheet = workbook.active
    return {"sheet_name": sheet.title, "physical_rows": sheet.max_row, "physical_columns": sheet.max_column,
            "second_section_header": list(next(sheet.iter_rows(min_row=35, max_row=35, values_only=True)))}


def main() -> None:
    results = {"run_id": "EXP021-SOURCE-003", "source_release_selection": "Drugs@FDA 2026-10-02 live release; Orange current ZIP last modified 2026-09-11; Purple latest listed August 2026 CSV and XLSX full-year dates", "sources": {}}
    for name, url in SOURCES.items():
        blob, meta = download(url)
        if name == "purple_book_august_2026":
            meta.update(inspect_purple(blob))
        elif name == "purple_book_august_2026_xlsx":
            meta.update(inspect_purple_xlsx(blob))
        else:
            meta["members"] = inspect_zip(blob, "\t" if name == "drugs_at_fda" else "~")
        results["sources"][name] = meta
    out = ROOT / "SOURCE_INDEX.json"
    out.write_text(json.dumps(results, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({key: {"bytes_read": value["bytes_read"], "members": {k: v["rows_after_header"] for k, v in value.get("members", {}).items()}, "physical_rows": value.get("physical_rows")} for key, value in results["sources"].items()}, indent=2))


if __name__ == "__main__":
    main()
