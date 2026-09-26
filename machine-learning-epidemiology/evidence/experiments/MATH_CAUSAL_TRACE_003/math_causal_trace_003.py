import torch, itertools, numpy as np, pandas as pd, matplotlib.pyplot as plt, math, shutil, zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path('/mnt/data')
OUT=BASE/'MATH_CAUSAL_TRACE_003'
OUT.mkdir(exist_ok=True)
torch.set_default_dtype(torch.float64)

# ---------------- model ----------------
BIT={0:-1.0,1:1.0}; AND_MARK=0.25; GIVES=-0.25; ANSWER=0.75
records=[]
for a,b in itertools.product([0,1], repeat=2):
    y=a & b
    q=[BIT[a],BIT[b]]
    cot=[AND_MARK, BIT[a], BIT[b], GIVES, BIT[y], ANSWER, BIT[y]]
    records.append((q,cot,y,(a,b)))

param_names=['W11','W12','W21','W22','u1','u2','b1','b2','v1','v2']
train_token_names=['Q_A','Q_B','AND','COT_A','COT_B','GIVES','RESULT','ANSWER','FINAL']
probe_token_names=['Q_A=1','Q_B=0','AND','A=1','B=0','GIVES','r=0','ANSWER','0']
eta=0.05; epochs=20

g=torch.Generator().manual_seed(1)
init32=(torch.randn(10,generator=g,dtype=torch.float32)*0.3)
init=init32.double()

def unpack(theta):
    return theta[:4].reshape(2,2), theta[4:6], theta[6:8], theta[8:10]

def run(theta,xs):
    W,u,b,v=unpack(theta); h=torch.zeros(2,dtype=theta.dtype); zs=[]; hs=[]
    for x in xs:
        x=torch.as_tensor(x,dtype=theta.dtype)
        h=torch.tanh(W@h+u*x+b); zs.append(v@h); hs.append(h)
    return torch.stack(zs),torch.stack(hs)

def run_preact(theta,xs):
    W,u,b,v=unpack(theta); h=torch.zeros(2,dtype=theta.dtype); zs=[]; hs=[]; acts=[]
    for x in xs:
        x=torch.as_tensor(x,dtype=theta.dtype)
        a=W@h+u*x+b; h=torch.tanh(a); acts.append(a); zs.append(v@h); hs.append(h)
    return torch.stack(zs),torch.stack(hs),torch.stack(acts)

def loss_theta_x(theta,xvec,y):
    z,h=run(theta,xvec); target=torch.tensor(float(y),dtype=theta.dtype)
    return torch.nn.functional.binary_cross_entropy_with_logits(z[1],target)+torch.nn.functional.binary_cross_entropy_with_logits(z[-1],target)

def loss_idx(theta,idx):
    q,cot,y,bits=records[idx]
    return loss_theta_x(theta,torch.tensor(q+cot,dtype=theta.dtype),y)

def update(th,idx):
    th=th.clone().detach().requires_grad_(True); L=loss_idx(th,idx); grad=torch.autograd.grad(L,th)[0]
    return (th-eta*grad).detach()

def grad_hess(th,idx):
    t=th.clone().detach().requires_grad_(True); L=loss_idx(t,idx)
    grad=torch.autograd.grad(L,t,create_graph=True)[0]
    H=torch.autograd.functional.hessian(lambda x:loss_idx(x,idx),t)
    return grad.detach(),H.detach()

def probe_logit(theta,k):
    q,cot,y,bits=records[2]
    z,h=run(theta,q+cot[:k]); return z[-1]

def train_order(order):
    th=init.clone()
    for ep in range(epochs):
        for idx in order: th=update(th,idx)
    return th

base=train_order([0,1,2,3]); swap=train_order([1,0,2,3])

# ---------------- 1) SGD commutator ----------------
comm_rows=[]; epoch_starts=[]
th=init.clone()
for ep in range(1,epochs+1):
    epoch_starts.append(th.clone())
    for idx in [0,1,2,3]: th=update(th,idx)
baseline_final=th.clone()

for ep,theta0 in enumerate(epoch_starts,1):
    AB=update(update(theta0,0),1); BA=update(update(theta0,1),0)
    exact=AB-BA
    ga,Ha=grad_hess(theta0,0); gb,Hb=grad_hess(theta0,1)
    approx=eta**2*(Hb@ga-Ha@gb)
    comm_rows.append({
        'epoch':ep,'exact_pair_delta_L2':float(torch.linalg.norm(exact)),
        'commutator_pred_L2':float(torch.linalg.norm(approx)),
        'cosine_exact_vs_prediction':float(torch.nn.functional.cosine_similarity(exact,approx,dim=0)),
        'relative_L2_error':float(torch.linalg.norm(exact-approx)/torch.linalg.norm(exact))
    })
comm_df=pd.DataFrame(comm_rows)
comm_df.to_csv(OUT/'order_commutator_by_epoch.csv',index=False)

# epoch 1 per-parameter vector
AB=update(update(init,0),1); BA=update(update(init,1),0); de=AB-BA
ga,Ha=grad_hess(init,0); gb,Hb=grad_hess(init,1); dap=eta**2*(Hb@ga-Ha@gb)
comm_param=pd.DataFrame({'parameter':param_names,'exact_AB_minus_BA':de.numpy(),'eta2_commutator_prediction':dap.numpy(),'raw_bracket':((Hb@ga-Ha@gb)).numpy()})
comm_param.to_csv(OUT/'order_commutator_parameter_vector_epoch1.csv',index=False)

# ---------------- 2) training token -> parameter -> CoT ----------------
def token_param_transfer(theta_pre,rec_idx):
    q,cot,y,bits=records[rec_idx]
    x0=torch.tensor(q+cot,dtype=torch.float64,requires_grad=True)
    th=theta_pre.clone().detach().requires_grad_(True)
    def post_from_x(x):
        L=loss_theta_x(th,x,y); gtheta=torch.autograd.grad(L,th,create_graph=True)[0]
        return th-eta*gtheta
    thpost=post_from_x(x0)
    P=torch.autograd.functional.jacobian(post_from_x,x0).detach() # 10x9
    return thpost.detach(),P

token_rows=[]; transfer_rows=[]; closure_errors=[]
for idx,(q,cot,y,bits) in enumerate(records):
    post,P=token_param_transfer(init,idx)
    J=[]
    for k in range(8):
        t=post.clone().requires_grad_(True); z=probe_logit(t,k); J.append(torch.autograd.grad(z,t)[0].detach())
    J=torch.stack(J); K=J@P
    # direct autograd closure
    x=torch.tensor(q+cot,dtype=torch.float64,requires_grad=True); th0=init.clone().detach().requires_grad_(True)
    L=loss_theta_x(th0,x,y); gth=torch.autograd.grad(L,th0,create_graph=True)[0]; post_graph=th0-eta*gth
    Kd=[]
    for k in range(8):
        z=probe_logit(post_graph,k); Kd.append(torch.autograd.grad(z,x,retain_graph=True)[0].detach())
    Kd=torch.stack(Kd); closure_errors.append(float(torch.max(torch.abs(K-Kd))))
    vals=q+cot
    for tp,name in enumerate(train_token_names):
        vec=P[:,tp].numpy(); top=int(np.argmax(np.abs(vec)))
        row={'record':str(bits),'target':y,'token_pos':tp,'token_name':name,'token_value':float(vals[tp]),
             'dtheta_dx_L2':float(np.linalg.norm(vec)),'top_parameter':param_names[top],'top_dtheta_dx':float(vec[top])}
        row.update({f'd_{pn}':float(vec[i]) for i,pn in enumerate(param_names)})
        token_rows.append(row)
        for k in range(8):
            transfer_rows.append({'record':str(bits),'target':y,'train_token_pos':tp,'train_token_name':name,
                                  'train_token_value':float(vals[tp]),'cot_k':k,
                                  'probe_prefix_last':'bare_Q' if k==0 else probe_token_names[k+1],
                                  'dz_probe_d_train_token':float(K[k,tp])})

token_df=pd.DataFrame(token_rows); transfer_df=pd.DataFrame(transfer_rows)
token_df.to_csv(OUT/'training_token_to_parameter_derivatives.csv',index=False)
transfer_df.to_csv(OUT/'training_token_to_cot_transfer.csv',index=False)

# ---------------- 3) exact final-history finite difference decomposition ----------------
q,cot,y,bits=records[2]; xs=q+cot
zA,hA,aA=run_preact(base,xs); zB,hB,aB=run_preact(swap,xs)
WA,uA,bA,vA=unpack(base); WB,uB,bB,vB=unpack(swap)
Wbar=(WA+WB)/2; vbar=(vA+vB)/2; dW=WA-WB; du=uA-uB; db=bA-bB; dv=vA-vB
channels=torch.zeros((8,2),dtype=torch.float64); exact_rows=[]
for t,x in enumerate(xs):
    hpA=torch.zeros(2) if t==0 else hA[t-1]; hpB=torch.zeros(2) if t==0 else hB[t-1]; hpbar=(hpA+hpB)/2
    da=aA[t]-aB[t]; dh=hA[t]-hB[t]
    S=torch.empty(2)
    for i in range(2):
        S[i]=dh[i]/da[i] if abs(float(da[i]))>1e-14 else 1-torch.tanh((aA[t,i]+aB[t,i])/2)**2
    new=torch.zeros_like(channels)
    for j in range(8): new[j]=S*(Wbar@channels[j])
    for p,(i,j) in enumerate([(0,0),(0,1),(1,0),(1,1)]):
        src=torch.zeros(2); src[i]=dW[i,j]*hpbar[j]; new[p]+=S*src
    for p,i in [(4,0),(5,1)]:
        src=torch.zeros(2); src[i]=du[i]*float(x); new[p]+=S*src
    for p,i in [(6,0),(7,1)]:
        src=torch.zeros(2); src[i]=db[i]; new[p]+=S*src
    channels=new
    contrib=[float(vbar@channels[j]) for j in range(8)]
    hbar=(hA[t]+hB[t])/2
    contrib += [float(dv[0]*hbar[0]),float(dv[1]*hbar[1])]
    row={'sequence_pos':t,'token':probe_token_names[t], 'coldstart_k':None if t==0 else t-1,
         'delta_h1':float(dh[0]),'delta_h2':float(dh[1]),'delta_z_base_minus_swap':float(zA[t]-zB[t]),
         'sum_parameter_contributions':float(sum(contrib))}
    row.update({f'contrib_{pn}':contrib[i] for i,pn in enumerate(param_names)})
    exact_rows.append(row)
exact_df=pd.DataFrame(exact_rows)
exact_df.to_csv(OUT/'exact_parameter_to_cot_decomposition.csv',index=False)
exact_z_resid=float(np.max(np.abs(exact_df['delta_z_base_minus_swap']-exact_df['sum_parameter_contributions'])))

# ---------------- 4) CoT tangent map ----------------
def cot_tangent(theta):
    x=torch.tensor(xs,dtype=torch.float64,requires_grad=True); z,h=run(theta,x); rows=[]
    for t in range(len(xs)): rows.append(torch.autograd.grad(z[t],x,retain_graph=True)[0].detach())
    return torch.stack(rows)

def cot_tangent_analytic(theta):
    W,u,b,v=unpack(theta); z,h,a=run_preact(theta,xs); T=len(xs); M=torch.zeros((T,T))
    A=[]; inp=[]
    for t in range(T):
        D=torch.diag(1-h[t]**2); A.append(D@W); inp.append(D@u)
    for t in range(T):
        for j in range(t+1):
            vec=inp[j].clone()
            for s in range(j+1,t+1): vec=A[s]@vec
            M[t,j]=v@vec
    return M
Tb=cot_tangent(base); Ts=cot_tangent(swap); Tba=cot_tangent_analytic(base); Tsa=cot_tangent_analytic(swap)
analytic_tangent_error=max(float(torch.max(torch.abs(Tb-Tba))),float(torch.max(torch.abs(Ts-Tsa))))
tangent_delta=Tb-Ts
for name,M in [('baseline',Tb),('swapped',Ts),('delta_base_minus_swap',tangent_delta)]:
    pd.DataFrame(M.numpy(),index=probe_token_names,columns=probe_token_names).to_csv(OUT/f'cot_tangent_{name}.csv')

# ---------------- 5) full causal propagation: local order event -> future training -> final CoT ----------------
schedule=[]
for ep in range(1,epochs+1):
    for pos,idx in enumerate([0,1,2,3]): schedule.append((ep,pos,idx))
pre_states=[]; update_J=[]; th=init.clone(); I=torch.eye(10)
for ep,pos,idx in schedule:
    pre_states.append(th.clone())
    t=th.clone().requires_grad_(True); H=torch.autograd.functional.hessian(lambda x:loss_idx(x,idx),t).detach()
    update_J.append(I-eta*H); th=update(th,idx)
baseline_final=th.clone()
Jfinal=[]
for k in range(8):
    t=baseline_final.clone().requires_grad_(True); z=probe_logit(t,k); Jfinal.append(torch.autograd.grad(z,t)[0].detach())
Jfinal=torch.stack(Jfinal)

def train_single_swap(ep_swap):
    th=init.clone()
    for ep in range(1,epochs+1):
        order=[1,0,2,3] if ep==ep_swap else [0,1,2,3]
        for idx in order: th=update(th,idx)
    return th

prop_rows=[]
for ep in range(1,epochs+1):
    start=(ep-1)*4; theta0=pre_states[start]
    AB=update(update(theta0,0),1); BA=update(update(theta0,1),0); dlocal=BA-AB
    ga,Ha=grad_hess(theta0,0); gb,Hb=grad_hess(theta0,1); dcomm=eta**2*(Ha@gb-Hb@ga) # swap - base
    Phi=I.clone()
    for j in range(start+2,len(schedule)): Phi=update_J[j]@Phi
    actual_th=train_single_swap(ep)
    for k in range(8):
        actual=float(probe_logit(actual_th,k)-probe_logit(baseline_final,k))
        p1=float(Jfinal[k]@(Phi@dlocal)); p2=float(Jfinal[k]@(Phi@dcomm))
        prop_rows.append({'epoch':ep,'cot_k':k,'actual_final_dz_swap_minus_base':actual,
                          'prediction_exact_local_delta_plus_tangent':p1,
                          'prediction_commutator_plus_tangent':p2})
prop_df=pd.DataFrame(prop_rows); prop_df.to_csv(OUT/'epoch_event_to_final_cot_causal_trace.csv',index=False)

def metrics(pred_col):
    a=prop_df['actual_final_dz_swap_minus_base'].to_numpy(); p=prop_df[pred_col].to_numpy()
    return {'corr':float(np.corrcoef(a,p)[0,1]),'rmse':float(np.sqrt(np.mean((a-p)**2))),
            'relative_L2_error':float(np.linalg.norm(a-p)/np.linalg.norm(a))}
metrics_exact=metrics('prediction_exact_local_delta_plus_tangent'); metrics_comm=metrics('prediction_commutator_plus_tangent')

# Sum over epoch events vs all-swapped history
sum_rows=[]
for k in range(8):
    sub=prop_df[prop_df.cot_k==k]
    actual=float(probe_logit(swap,k)-probe_logit(base,k))
    sum_rows.append({'cot_k':k,'actual_all_swapped_dz':actual,
                     'sum_actual_single_epoch_effects':float(sub.actual_final_dz_swap_minus_base.sum()),
                     'sum_exactlocal_tangent_predictions':float(sub.prediction_exact_local_delta_plus_tangent.sum()),
                     'sum_commutator_tangent_predictions':float(sub.prediction_commutator_plus_tangent.sum())})
sum_df=pd.DataFrame(sum_rows); sum_df.to_csv(OUT/'summed_history_prediction_by_cot.csv',index=False)

# ---------------- figures ----------------
# token -> parameter heatmap (all records stacked)
mat=token_df[[f'd_{p}' for p in param_names]].to_numpy()
fig,ax=plt.subplots(figsize=(10,9)); im=ax.imshow(mat,aspect='auto')
ax.set_xticks(range(10)); ax.set_xticklabels(param_names,rotation=45,ha='right')
labels=[f"{r.record}:{r.token_name}" for r in token_df.itertuples()]
ax.set_yticks(range(len(labels))); ax.set_yticklabels(labels,fontsize=6)
ax.set_title('Training-token derivative of one SGD parameter update')
fig.colorbar(im,ax=ax,label='d theta+ / d token encoding')
fig.tight_layout(); fig.savefig(OUT/'training_token_to_parameter_heatmap.png',dpi=180,bbox_inches='tight'); plt.close(fig)

# swapped pair training token -> CoT transfer heatmap
pair_transfer=transfer_df[transfer_df.record.isin(['(0, 0)','(0, 1)'])].copy()
rows_lbl=[]; M=[]
for rec in ['(0, 0)','(0, 1)']:
    for tp,name in enumerate(train_token_names):
        sub=pair_transfer[(pair_transfer.record==rec)&(pair_transfer.train_token_pos==tp)].sort_values('cot_k')
        rows_lbl.append(f'{rec}:{name}'); M.append(sub.dz_probe_d_train_token.to_numpy())
M=np.array(M)
fig,ax=plt.subplots(figsize=(9,7)); im=ax.imshow(M,aspect='auto')
ax.set_xticks(range(8)); ax.set_xticklabels([f'k={k}' for k in range(8)])
ax.set_yticks(range(len(rows_lbl))); ax.set_yticklabels(rows_lbl,fontsize=7)
ax.set_title('Training-token -> final CoT-slice transfer after one SGD update')
fig.colorbar(im,ax=ax,label='d z_k / d training-token encoding')
fig.tight_layout(); fig.savefig(OUT/'training_token_to_cot_heatmap.png',dpi=180,bbox_inches='tight'); plt.close(fig)

# exact per-parameter CoT decomposition heatmap
cold=exact_df[exact_df.sequence_pos>=1].copy(); MC=np.array([[row[f'contrib_{p}'] for row in cold.to_dict('records')] for p in param_names])
fig,ax=plt.subplots(figsize=(9,6)); im=ax.imshow(MC,aspect='auto')
ax.set_xticks(range(8)); ax.set_xticklabels([f'k={k}' for k in range(8)])
ax.set_yticks(range(10)); ax.set_yticklabels(param_names)
ax.set_title('Exact contribution of each parameter-history difference to CoT logit difference')
fig.colorbar(im,ax=ax,label='exact contribution to delta z')
fig.tight_layout(); fig.savefig(OUT/'exact_parameter_to_cot_heatmap.png',dpi=180,bbox_inches='tight'); plt.close(fig)

# epoch causal prediction at k=2
k2=prop_df[prop_df.cot_k==2].sort_values('epoch')
fig,ax=plt.subplots(figsize=(10,5)); ax.plot(k2.epoch,k2.actual_final_dz_swap_minus_base,marker='o',label='actual single-epoch intervention')
ax.plot(k2.epoch,k2.prediction_commutator_plus_tangent,marker='o',label='commutator + tangent prediction')
ax.set_xlabel('epoch containing the local order swap'); ax.set_ylabel('final k=2 logit effect')
ax.set_title('Mathematical prediction of each historical event at the final CoT slice'); ax.legend(); fig.tight_layout()
fig.savefig(OUT/'epoch_to_final_k2_prediction.png',dpi=180,bbox_inches='tight'); plt.close(fig)

# CoT tangent delta
fig,ax=plt.subplots(figsize=(8,7)); im=ax.imshow(tangent_delta.numpy(),aspect='auto')
ax.set_xticks(range(9)); ax.set_xticklabels(probe_token_names,rotation=45,ha='right')
ax.set_yticks(range(9)); ax.set_yticklabels(probe_token_names)
ax.set_title('Difference between the two history-conditioned CoT tangent maps')
fig.colorbar(im,ax=ax,label='delta (d z_t / d x_j)')
fig.tight_layout(); fig.savefig(OUT/'cot_tangent_map_difference.png',dpi=180,bbox_inches='tight'); plt.close(fig)

# ---------------- report ----------------
comm0=comm_df.iloc[0]
report=OUT/'MATH_CAUSAL_TRACE_003.md'
with report.open('w',encoding='utf-8') as f:
    f.write('# MATH_CAUSAL_TRACE_003 | Mathematical causal tracing of training history\n\n')
    f.write('## Object\n\nSame 10-parameter recurrent network and local order intervention used in MICRO-001/CAUSAL-002. Derivative audit is evaluated in float64 from the same seeded initial parameter values to suppress numerical noise.\n\n')
    f.write('## 1. Local training-order effect is an SGD commutator\n\n')
    f.write('For U_s(theta)=theta-eta g_s(theta), the adjacent-order difference obeys\n\n')
    f.write('theta_AB - theta_BA = eta^2 [H_b g_a - H_a g_b] + O(eta^3).\n\n')
    f.write(f"Epoch 1: exact norm={comm0.exact_pair_delta_L2:.9f}; predicted norm={comm0.commutator_pred_L2:.9f}; cosine={comm0.cosine_exact_vs_prediction:.6f}; relative error={100*comm0.relative_L2_error:.3f}%.\n\n")
    f.write(f"Across all 20 baseline epoch-start states: mean cosine={comm_df.cosine_exact_vs_prediction.mean():.6f}; minimum cosine={comm_df.cosine_exact_vs_prediction.min():.6f}; mean relative error={100*comm_df.relative_L2_error.mean():.3f}%.\n\n")
    f.write('## 2. Training token -> parameter update -> CoT slice\n\n')
    f.write('Treating the fixed scalar token encoding x_t as the intervention coordinate, one SGD step gives the exact local derivative\n\n')
    f.write('d theta+ / d x_t = -eta * d^2 L / (d theta d x_t).\n\n')
    f.write('For CoT slice k, J_k = d z_k / d theta, therefore\n\n')
    f.write('d z_k / d x_t = J_k (d theta+ / d x_t).\n\n')
    f.write(f"Direct autograd and the matrix product J_k P agree to max absolute error {max(closure_errors):.3e}.\n\n")
    top=token_df.sort_values('dtheta_dx_L2',ascending=False).head(10)[['record','token_name','token_value','dtheta_dx_L2','top_parameter','top_dtheta_dx']]
    f.write('Strongest local token-to-parameter derivatives at initialization:\n\n'); f.write(top.to_markdown(index=False)); f.write('\n\n')
    f.write('## 3. Exact finite-difference decomposition over CoT\n\n')
    f.write('For the two final history-conditioned models, with bar denoting midpoint and Delta denoting baseline minus swapped:\n\n')
    f.write('Delta h_t = S_t [ Wbar Delta h_(t-1) + Delta W hbar_(t-1) + Delta u x_t + Delta b ],\n\n')
    f.write('Delta z_t = vbar^T Delta h_t + Delta v^T hbar_t.\n\n')
    f.write('S_t is the elementwise secant slope of tanh between the two trajectories. The additive source terms were recursively split into all 10 parameter channels.\n\n')
    f.write(f"Maximum reconstruction residual for Delta z across the probe sequence: {exact_z_resid:.3e}.\n\n")
    f.write(cold[['coldstart_k','token','delta_z_base_minus_swap']+[f'contrib_{p}' for p in param_names]].to_markdown(index=False)); f.write('\n\n')
    f.write('## 4. Training event -> future training -> final CoT\n\n')
    f.write('For a local swap event in epoch e, later SGD updates have tangent map A_s = I - eta H_s. Let Phi be their ordered product. Then\n\n')
    f.write('Delta z_(k,T)^(e) ~= J_(k,T) Phi_(T<-e) Delta theta_e,\n\n')
    f.write('and substituting the local commutator gives a closed history-to-CoT approximation.\n\n')
    f.write(f"Using exact local swap Delta theta plus tangent propagation over 20 epochs x 8 CoT slices: correlation={metrics_exact['corr']:.6f}, relative L2 error={100*metrics_exact['relative_L2_error']:.3f}%.\n\n")
    f.write(f"Using only the Hessian/gradient commutator plus tangent propagation: correlation={metrics_comm['corr']:.6f}, relative L2 error={100*metrics_comm['relative_L2_error']:.3f}%.\n\n")
    f.write('## 5. CoT itself has a history-conditioned tangent geometry\n\n')
    f.write('For D_t=diag(1-h_t^2), A_t=D_t W, the local influence of an earlier input token x_j on later logit z_t is\n\n')
    f.write('d z_t / d x_j = v^T A_t A_(t-1) ... A_(j+1) D_j u,  j<=t.\n\n')
    f.write(f"Analytic recurrence and autograd agree to max absolute error {analytic_tangent_error:.3e}. The Frobenius norm of the difference between the baseline and swapped CoT tangent maps is {float(torch.linalg.norm(tangent_delta)):.6f}.\n\n")
    f.write('## Evidence statement\n\n')
    f.write('In this controlled 10-parameter system, the local temporal order of training examples produces a parameter displacement captured by the noncommuting SGD update fields. That displacement is propagated by the tangent dynamics of later training and is read out differently across CoT prefixes. Training-token, parameter, hidden-state, and CoT-output effects can therefore be connected by explicit derivative and finite-difference identities in this system.\n')

# ---------------- update Living Record DOCX ----------------
source=BASE/'ML_Epidemiology_Case_Dissection_Living_Log_v0.1_20260926.docx'
dest=BASE/'ML_Epidemiology_Case_Dissection_Living_Log_v0.2_20260926.docx'
doc=Document(source)

# Helpers using existing style
NAVY='17365D'; BLUE='2F5597'; LIGHT='EAF0F8'; GREEN='375623'; GREEN_BG='E2F0D9'; MUTED='5F6368'; DARK='202124'
def set_font(run,size=10,bold=None,color=None,italic=None,name='Noto Sans CJK SC'):
    run.font.name=name; rpr=run._element.get_or_add_rPr(); rpr.rFonts.set(qn('w:ascii'),name); rpr.rFonts.set(qn('w:hAnsi'),name); rpr.rFonts.set(qn('w:eastAsia'),name)
    run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color: run.font.color.rgb=RGBColor.from_string(color)
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def add_callout(title,body):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); shade(c,GREEN_BG)
    p=c.paragraphs[0]; r=p.add_run(title+'\n'); set_font(r,10.5,True,GREEN); r=p.add_run(body); set_font(r,9.5,color=DARK)
def add_table(df,cols,headers=None,widths=None,fs=7.7,dec=6):
    headers=headers or cols; t=doc.add_table(rows=1,cols=len(cols)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h); shade(t.rows[0].cells[i],NAVY)
        for r in t.rows[0].cells[i].paragraphs[0].runs: set_font(r,fs,True,'FFFFFF')
    for _,row in df.iterrows():
        cells=t.add_row().cells
        for i,cname in enumerate(cols):
            val=row[cname]
            if isinstance(val,(float,np.floating)): val=f'{float(val):.{dec}f}'
            cells[i].text=str(val)
            for r in cells[i].paragraphs[0].runs: set_font(r,fs)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    return t

# replace visible version strings and current series metadata
for p in list(doc.paragraphs)+[p for sec in doc.sections for p in list(sec.header.paragraphs)+list(sec.footer.paragraphs)]:
    if 'v0.1' in p.text:
        for r in p.runs:
            if 'v0.1' in r.text: r.text=r.text.replace('v0.1','v0.2')
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            txt=cell.text
            if txt=='v0.1': cell.text='v0.2'
            elif 'TEN_PARAM_NN_MICROSCOPE_001 / TEN_PARAM_HISTORY_CAUSAL_002' in txt:
                cell.text='TEN_PARAM_NN_MICROSCOPE_001 / TEN_PARAM_HISTORY_CAUSAL_002 / MATH_CAUSAL_TRACE_003'

# append section
doc.add_page_break(); doc.add_heading('3. MATH-003 | 训练历史的数学因果追溯',level=1)
add_callout('实验问题','把“训练顺序 → 参数历史 → CoT 条件响应”从现象链推进为可计算方程：训练 token 怎样改变参数更新；相邻样本顺序怎样产生非交换位移；这些位移怎样经过后续训练传播，并最终映射到每个 CoT-prefix 横截面。')

p=doc.add_paragraph(); r=p.add_run('核心闭合式：'); set_font(r,10,True,NAVY); r=p.add_run(' Δz(k,T,e) ≈ J(k,T) · Φ(T←e) · η²(Ha·gb − Hb·ga) '); set_font(r,10)
p=doc.add_paragraph(); r=p.add_run('其中 J 是最终 CoT 横截面对参数的 Jacobian，Φ 是后续 SGD 更新 Jacobian (I−ηH) 的有序乘积，最右侧是局部顺序交换的二阶非交换项。'); set_font(r,9.5)

doc.add_heading('3.1 局部训练顺序 = 非交换的 SGD 更新场',level=2)
p=doc.add_paragraph(); r=p.add_run('对 Ua(θ)=θ−ηga(θ)、Ub(θ)=θ−ηgb(θ)，相邻样本交换满足：'); set_font(r,10)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('θAB − θBA = η²(Hb ga − Ha gb) + O(η³)'); set_font(r,11,True,BLUE)
add_table(comm_df.head(6),['epoch','exact_pair_delta_L2','commutator_pred_L2','cosine_exact_vs_prediction','relative_L2_error'],
          headers=['epoch','||Δθ|| exact','||Δθ|| formula','cosine','relative error'],widths=[0.6,1.25,1.25,1.0,1.1],fs=7.6)
p=doc.add_paragraph(); r=p.add_run(f"Epoch 1：方向余弦 {comm0.cosine_exact_vs_prediction:.6f}，相对误差 {100*comm0.relative_L2_error:.2f}%。20 个 epoch-start 状态中平均方向余弦 {comm_df.cosine_exact_vs_prediction.mean():.6f}，最小值 {comm_df.cosine_exact_vs_prediction.min():.6f}。"); set_font(r,9.5)

doc.add_heading('3.2 训练 token → 参数 → CoT 的链式导数',level=2)
p=doc.add_paragraph(); r=p.add_run('固定 token 采用标量编码 x_t。对一次 SGD 更新：'); set_font(r,10)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('∂θ⁺/∂x_t = −η · ∂²L/(∂θ∂x_t)'); set_font(r,11,True,BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('∂z_k/∂x_t = (∂z_k/∂θ)(∂θ⁺/∂x_t)'); set_font(r,11,True,BLUE)
p=doc.add_paragraph(); r=p.add_run(f"直接自动微分与矩阵链乘 J·P 的最大绝对闭合误差为 {max(closure_errors):.3e}。这张 transfer matrix 直接把训练 token 的局部数值扰动映射到后续 CoT 第 k 层的 logit。 "); set_font(r,9.5)
doc.add_picture(str(OUT/'training_token_to_cot_heatmap.png'),width=Inches(6.4))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('图 3.1 训练 token → 一次参数更新 → CoT-prefix logit 的局部传递矩阵。'); set_font(r,8.3,color=MUTED,italic=True)

doc.add_heading('3.3 两条最终历史在 CoT 上的精确有限差分解剖',level=2)
p=doc.add_paragraph(); r=p.add_run('对 baseline 与 swapped 两只最终模型，同一 token x_t 下的隐藏态差满足精确恒等式：'); set_font(r,10)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Δh_t = S_t[ W̄Δh_(t−1) + ΔW h̄_(t−1) + Δu x_t + Δb ]'); set_font(r,10.5,True,BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Δz_t = v̄ᵀΔh_t + Δvᵀh̄_t'); set_font(r,10.5,True,BLUE)
p=doc.add_paragraph(); r=p.add_run(f"S_t 为 tanh 在两条轨迹之间的逐维割线斜率。按 10 个参数分别递推后，各参数贡献之和重建真实 Δz 的最大残差为 {exact_z_resid:.3e}。 "); set_font(r,9.5)
doc.add_picture(str(OUT/'exact_parameter_to_cot_heatmap.png'),width=Inches(6.35))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('图 3.2 每块“参数砖”对不同 CoT 横截面 logit 差的精确贡献。'); set_font(r,8.3,color=MUTED,italic=True)

doc.add_heading('3.4 历史事件 → 后续训练 → 最终 CoT',level=2)
p=doc.add_paragraph(); r=p.add_run('每个后续 SGD update 的局部传播矩阵为 A_s=I−ηH_s。把局部 swap 位移送入所有未来 update，再由最终 CoT Jacobian 读出：'); set_font(r,10)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Δz(k,T,e) ≈ J(k,T) · [Π A_s] · Δθ_e'); set_font(r,11,True,BLUE)
metricdf=pd.DataFrame([
    ['真实局部 swap Δθ + tangent propagation',metrics_exact['corr'],metrics_exact['relative_L2_error']*100],
    ['Hessian/gradient commutator + tangent propagation',metrics_comm['corr'],metrics_comm['relative_L2_error']*100],
],columns=['预测链','160 个观测点相关','relative L2 error (%)'])
add_table(metricdf,['预测链','160 个观测点相关','relative L2 error (%)'],widths=[3.3,1.3,1.4],fs=8.4)
doc.add_picture(str(OUT/'epoch_to_final_k2_prediction.png'),width=Inches(6.35))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('图 3.3 只交换某一个 epoch 的局部顺序：数学传播式预测其最终 k=2 CoT 效应。'); set_font(r,8.3,color=MUTED,italic=True)

doc.add_heading('3.5 CoT 本身也具有历史条件化的切向几何',level=2)
p=doc.add_paragraph(); r=p.add_run('令 D_t=diag(1−h_t²)、A_t=D_tW，则早期 token x_j 对后续 z_t 的局部影响为：'); set_font(r,10)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('∂z_t/∂x_j = vᵀ A_t A_(t−1)…A_(j+1) D_j u'); set_font(r,10.5,True,BLUE)
p=doc.add_paragraph(); r=p.add_run(f"解析递推与 autograd 的最大误差 {analytic_tangent_error:.3e}。两条训练历史最终形成的 CoT tangent map Frobenius 距离为 {float(torch.linalg.norm(tangent_delta)):.6f}。同一串 token 因而在两块历史地基上具有不同的局部传播几何。 "); set_font(r,9.5)
doc.add_picture(str(OUT/'cot_tangent_map_difference.png'),width=Inches(6.15))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('图 3.4 两条训练历史下 CoT token-to-token tangent map 的差。'); set_font(r,8.3,color=MUTED,italic=True)

add_callout('MATH-003 当前证据链','在这只 10 参数可完全观测系统中，局部训练顺序的影响可以由 SGD 梯度场的非交换项定量给出；训练 token 到参数的一次更新由混合二阶导数给出；后续训练通过 Hessian tangent map 传播这一扰动；最终 CoT-prefix logit 再由参数 Jacobian 读出。两条最终历史之间的逐 token 状态/输出差还可以用精确有限差分递推逐参数重建。')

doc.add_heading('3.6 MATH-003 原始数据索引',level=2)
idxdf=pd.DataFrame([
    ['order_commutator_by_epoch.csv','20 个 epoch-start 的真实顺序差与 Hessian/gradient 非交换预测'],
    ['training_token_to_parameter_derivatives.csv','36 个训练 token 对 10 参数一次 SGD 更新的混合偏导'],
    ['training_token_to_cot_transfer.csv','训练 token → 参数 → 8 个 CoT 横截面的链式 transfer'],
    ['exact_parameter_to_cot_decomposition.csv','两条最终历史的 10 参数逐 CoT 精确有限差分解剖'],
    ['epoch_event_to_final_cot_causal_trace.csv','20 个局部历史事件 → 后续训练 → 8 个最终 CoT 横截面'],
    ['cot_tangent_delta_base_minus_swap.csv','两条历史条件化 CoT tangent map 的差'],
    ['MATH_CAUSAL_TRACE_003.md','本实验数学说明与关键结果'],
],columns=['文件','内容'])
add_table(idxdf,['文件','内容'],widths=[2.5,4.5],fs=7.8)

# version history table at end
doc.add_heading('版本更新',level=1)
ver=pd.DataFrame([['v0.2','2026-09-26','新增 MATH-003：SGD 非交换子、训练 token 混合导数、tangent propagation、逐参数 CoT 精确有限差分解剖。']],columns=['版本','日期','新增内容'])
add_table(ver,['版本','日期','新增内容'],widths=[0.7,1.2,5.6],fs=8.2)

doc.save(dest)

# zip experiment package (exclude docx, living record separate)
zip_path=BASE/'MATH_CAUSAL_TRACE_003.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as zf:
    for p in sorted(OUT.iterdir()): zf.write(p,arcname=p.name)
    zf.write(BASE/'math_causal_trace_003.py',arcname='math_causal_trace_003.py')

print('OUTPUT_DIR',OUT)
print('REPORT',report)
print('DOCX',dest)
print('ZIP',zip_path)
print('COMM_EPOCH1',dict(comm0))
print('COMM_MEAN_COS',comm_df.cosine_exact_vs_prediction.mean())
print('CHAIN_CLOSURE',max(closure_errors))
print('EXACT_Z_RESID',exact_z_resid)
print('PROP_EXACT',metrics_exact)
print('PROP_COMM',metrics_comm)
print('TANGENT_ERR',analytic_tangent_error,'TANGENT_DELTA_NORM',float(torch.linalg.norm(tangent_delta)))
