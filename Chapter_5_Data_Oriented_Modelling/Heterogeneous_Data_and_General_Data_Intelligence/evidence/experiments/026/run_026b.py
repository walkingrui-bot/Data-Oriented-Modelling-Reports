import os
import re, math, time, json, hashlib, multiprocessing as mp
from pathlib import Path
from collections import defaultdict
import numpy as np, pandas as pd
from docx import Document
from scipy import sparse
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.decomposition import TruncatedSVD
from sklearn.cluster import KMeans
from sklearn.preprocessing import normalize

OUT_DEFAULT=Path(os.environ.get('B3_OUTPUT_DIR', str(Path.cwd()/'reproduction_026')))
OUT_DEFAULT.mkdir(exist_ok=True, parents=True)
DOC=Path(os.environ['B3_CORPUS_DOCX'])
OUT=Path(os.environ.get('B3_OUTPUT_DIR', str(Path.cwd()/'reproduction_026'))); TOPK=5; KS=[2,4,8,16]

def extract(path):
 d=Document(path); chunks=[]; cur=None; chap='Front matter'
 for idx,p in enumerate(d.paragraphs):
  t=p.text.strip(); style=p.style.name if p.style else ''
  if not t: continue
  h=style.startswith('Heading') or 'Title' in style
  if h:
   if cur: cur['text']='\n'.join(cur['parts']); chunks.append(cur)
   if style=='Heading 1': chap=t
   cur={'id':f'P{idx:04d}','kind':'section','heading':t,'chapter':chap,'parts':[t]}
  else:
   if cur is None: cur={'id':'FRONT','kind':'section','heading':'Front matter','chapter':'Front matter','parts':[]}
   cur['parts'].append(t)
 if cur: cur['text']='\n'.join(cur['parts']); chunks.append(cur)
 for ti,tbl in enumerate(d.tables):
  rows=[]
  for row in tbl.rows:
   vals=[c.text.strip().replace('\n',' ') for c in row.cells]
   if any(vals): rows.append(' | '.join(vals))
  txt='\n'.join(rows)
  if txt: chunks.append({'id':f'T{ti:03d}','kind':'table','heading':f'Table {ti+1}','chapter':'Tables','parts':rows,'text':txt})
 for c in chunks: c['bytes']=len(c['text'].encode())
 return chunks
chunks=extract(DOC); texts=[c['text'] for c in chunks]; N=len(chunks); barr=np.array([c['bytes'] for c in chunks]); raw=int(barr.sum())
vec=TfidfVectorizer(analyzer='char',ngram_range=(2,5),min_df=2,max_features=30000,norm='l2')
X=vec.fit_transform(texts).tocsr(); V=X.shape[1]
svd=TruncatedSVD(n_components=64,random_state=0); Y=normalize(svd.fit_transform(X))

def clean(s,n=180):
 s=re.sub(r'\s+',' ',s).strip(); s=re.sub(r'^\d+(?:\.\d+)*\.?\s*','',s); return s[:n]
queries=[]
for c in chunks:
 if c['kind']!='section': continue
 body=' '.join(c['text'].split('\n')[1:]).strip(); q=clean(body if len(body)>=60 else c['heading'])
 if len(q)>=20: queries.append(q)
secs=[c for c in chunks if c['kind']=='section']
for a,b in zip(secs[:-1],secs[1:]):
 if a['chapter']==b['chapter'] and a['chapter']!='Front matter':
  q=(clean(a['heading'],80)+' '+clean(b['heading'],80)).strip()
  if len(q)>=25: queries.append(q)
if len(queries)>500:
 ix=np.linspace(0,len(queries)-1,500,dtype=int); queries=[queries[i] for i in ix]
Q=vec.transform(queries).tocsr(); QY=normalize(svd.transform(Q)); qn=len(queries)
CS=(Q@X.T).toarray()
def topk(a,k=5):
 if len(a)<=k: return np.argsort(-a)
 i=np.argpartition(-a,k-1)[:k]; return i[np.argsort(-a[i])]
ref=[topk(CS[i]) for i in range(qn)]
S=Y@Y.T; np.fill_diagonal(S,-1); edges=set()
for i in range(N):
 for j in topk(S[i],5): edges.add(tuple(sorted((i,int(j)))))

rows=[]; route=[]; assigns={}
for k in KS:
 km=KMeans(n_clusters=k,random_state=0,n_init=20).fit(Y); centers=normalize(km.cluster_centers_)
 sims=Y@centers.T
 cap=raw/k*1.08
 order=np.argsort(-barr) # place large objects first
 load=np.zeros(k,float); a=np.full(N,-1,int)
 for i in order:
  prefs=np.argsort(-sims[i])
  chosen=None
  for s in prefs:
   if load[s]+barr[i] <= cap:
    chosen=int(s); break
  if chosen is None: chosen=int(np.argmin(load))
  a[i]=chosen; load[chosen]+=barr[i]
 assigns[k]=a
 count=np.bincount(a,minlength=k); cut=sum(a[i]!=a[j] for i,j in edges)/len(edges)
 rows.append({'policy':'balanced_svd_semantic','shards':k,'byte_load_cv':float(load.std()/load.mean()),'max_byte_share':float(load.max()/load.sum()),'chunk_count_cv':float(count.std()/count.mean()),'semantic_knn_edge_cut':float(cut),'target_cap_share':1.08/k})
 for p in [1,4]:
  prot=[]; proto_shard=[]
  for s in range(k):
   idx=np.where(a==s)[0]
   if len(idx)==0: continue
   pp=min(p,len(idx))
   if pp==1:
    c=normalize(Y[idx].mean(axis=0,keepdims=True))
   else:
    c=normalize(KMeans(n_clusters=pp,random_state=0,n_init=10).fit(Y[idx]).cluster_centers_)
   prot.append(c); proto_shard.extend([s]*len(c))
  P=np.vstack(prot); proto_shard=np.array(proto_shard)
  PS=QY@P.T
  shardscore=np.full((qn,k),-np.inf)
  for z,s in enumerate(proto_shard): shardscore[:,s]=np.maximum(shardscore[:,s],PS[:,z])
  for m0 in [1,2,4,8]:
   m=min(k,m0); rec=[]; ex=[]; scored=[]; spans=[]
   for qi in range(qn):
    ss=topk(shardscore[qi],m); cand=np.where(np.isin(a,ss))[0]; got=cand[topk(CS[qi,cand],TOPK)]
    R=set(map(int,ref[qi])); G=set(map(int,got)); rec.append(len(R&G)/TOPK); ex.append(R==G); scored.append(len(cand)/N); spans.append(len(set(a[list(R)])))
   route.append({'policy':'balanced_svd_semantic','shards':k,'prototypes_per_shard':p,'route_shards':m,'mean_recall_at5':float(np.mean(rec)),'exact_top5_set_rate':float(np.mean(ex)),'mean_scored_fraction':float(np.mean(scored)),'mean_reference_shard_span':float(np.mean(spans)),'router_summary_bytes':int(P.size*4),'summary_to_raw_ratio':float(P.size*4/raw)})
pd.DataFrame(rows).to_csv(OUT/'do026_balanced_sharding.csv',index=False)
pd.DataFrame(route).to_csv(OUT/'do026_compact_router.csv',index=False)
# save 8-shard assignment for process benchmark
np.save(OUT/'do026_balanced_assign8.npy',assigns[8]); sparse.save_npz(OUT/'do026_X.npz',X); sparse.save_npz(OUT/'do026_Q.npz',Q)
print(pd.DataFrame(rows).to_string(index=False))
print('\nROUTER 8/16:')
print(pd.DataFrame(route).query('shards>=8').to_string(index=False))
