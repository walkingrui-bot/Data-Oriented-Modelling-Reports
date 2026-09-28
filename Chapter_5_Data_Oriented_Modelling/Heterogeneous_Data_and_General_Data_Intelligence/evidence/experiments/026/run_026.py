import os
import os, re, json, math, time, hashlib, statistics, multiprocessing as mp
from pathlib import Path
from collections import defaultdict
import numpy as np
import pandas as pd
from docx import Document
from scipy import sparse
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.cluster import MiniBatchKMeans
from sklearn.preprocessing import normalize

DOC=Path(os.environ['B3_CORPUS_DOCX'])
OUT=Path(os.environ.get('B3_OUTPUT_DIR', str(Path.cwd()/'reproduction_026')))
OUT.mkdir(exist_ok=True, parents=True)
TOPK=5
KS=[2,4,8,16]
POLICIES=['hash','chapter_greedy','semantic_kmeans']

# ---------- Extract one real logical body ----------
def extract_chunks(path):
    d=Document(path)
    chunks=[]
    current=None
    chapter='Front matter'
    for idx,p in enumerate(d.paragraphs):
        txt=p.text.strip()
        if not txt: continue
        style=p.style.name if p.style else ''
        ishead=style.startswith('Heading') or 'Title' in style
        if ishead:
            if current is not None:
                current['text']='\n'.join(current['parts']).strip()
                chunks.append(current)
            if style=='Heading 1': chapter=txt
            current={'id':f'P{idx:04d}','kind':'section','heading':txt,'chapter':chapter,'parts':[txt], 'para_index':idx}
        else:
            if current is None:
                current={'id':'FRONT','kind':'section','heading':'Front matter','chapter':'Front matter','parts':[], 'para_index':-1}
            current['parts'].append(txt)
    if current is not None:
        current['text']='\n'.join(current['parts']).strip(); chunks.append(current)
    # tables as state objects, attached to closest preceding chapter context by document order unavailable;
    # use their own contents and sequential id. This preserves all table text as distributed objects.
    for ti,tbl in enumerate(d.tables):
        rows=[]
        for row in tbl.rows:
            vals=[c.text.strip().replace('\n',' ') for c in row.cells]
            if any(vals): rows.append(' | '.join(vals))
        text='\n'.join(rows).strip()
        if text:
            chunks.append({'id':f'T{ti:03d}','kind':'table','heading':f'Table {ti+1}','chapter':'Tables','parts':rows,'text':text,'para_index':10**6+ti})
    for c in chunks:
        c['utf8_bytes']=len(c['text'].encode('utf-8'))
        c['chars']=len(c['text'])
    return chunks

chunks=extract_chunks(DOC)
texts=[c['text'] for c in chunks]
ids=[c['id'] for c in chunks]
id2i={x:i for i,x in enumerate(ids)}
raw_bytes=sum(c['utf8_bytes'] for c in chunks)

# Mixed Chinese/English: character n-grams avoid language-specific tokenization assumptions.
vec=TfidfVectorizer(analyzer='char', ngram_range=(2,5), min_df=2, max_features=30000, norm='l2', dtype=np.float64)
X=vec.fit_transform(texts).tocsr()
N,V=X.shape

# ---------- Deterministic real queries from the body itself ----------
def clean_snippet(s, max_chars=180):
    s=re.sub(r'\s+',' ',s).strip()
    # Remove leading experiment numbering so routing is not solved by a single literal id.
    s=re.sub(r'^\d+(?:\.\d+)*\.?\s*','',s)
    s=re.sub(r'^[A-Z0-9][A-Z0-9\-×²_]+\s*[—|:-]+\s*','',s)
    return s[:max_chars]

queries=[]; qmeta=[]
# one heading/body query per section where possible
for i,c in enumerate(chunks):
    if c['kind']!='section': continue
    body=' '.join(c['text'].split('\n')[1:]).strip()
    src=body if len(body)>=60 else c['heading']
    q=clean_snippet(src,180)
    if len(q)>=20:
        queries.append(q); qmeta.append({'type':'single','src':[c['id']]})
# composite queries from adjacent sections within same chapter
sections=[c for c in chunks if c['kind']=='section']
for a,b in zip(sections[:-1],sections[1:]):
    if a['chapter']==b['chapter'] and a['chapter']!='Front matter':
        qa=clean_snippet(a['heading'],80); qb=clean_snippet(b['heading'],80)
        q=(qa+' '+qb).strip()
        if len(q)>=25:
            queries.append(q); qmeta.append({'type':'composite','src':[a['id'],b['id']]})
# deterministic cap keeps runtime modest while spanning the entire record
if len(queries)>500:
    idx=np.linspace(0,len(queries)-1,500,dtype=int)
    queries=[queries[i] for i in idx]; qmeta=[qmeta[i] for i in idx]
Q=vec.transform(queries).tocsr()
Qn=len(queries)

# central exact reference
def topk_from_scores(arr,k=TOPK):
    if arr.size<=k:
        return np.argsort(-arr)
    ind=np.argpartition(-arr,k-1)[:k]
    return ind[np.argsort(-arr[ind])]

central_scores=(Q @ X.T).toarray()
central_top=[topk_from_scores(central_scores[i]) for i in range(Qn)]

# Semantic dependency graph: top-5 nearest other chunks in the actual ontology.
S=(X @ X.T).toarray(); np.fill_diagonal(S,-1)
knn_edges=set()
for i in range(N):
    js=topk_from_scores(S[i],5)
    for j in js:
        a,b=sorted((i,int(j)))
        knn_edges.add((a,b))

bytes_arr=np.array([c['utf8_bytes'] for c in chunks])

def assignments(policy,k):
    if policy=='hash':
        a=np.array([int(hashlib.sha1(cid.encode()).hexdigest(),16)%k for cid in ids],dtype=int)
    elif policy=='chapter_greedy':
        groups=defaultdict(list)
        for i,c in enumerate(chunks): groups[c['chapter']].append(i)
        group_items=sorted(groups.items(), key=lambda kv:-sum(bytes_arr[kv[1]]))
        load=[0]*k; a=np.empty(N,dtype=int)
        for g,idxs in group_items:
            s=min(range(k), key=lambda z:load[z])
            a[idxs]=s; load[s]+=int(bytes_arr[idxs].sum())
    elif policy=='semantic_kmeans':
        km=MiniBatchKMeans(n_clusters=k, random_state=0, n_init=10, batch_size=min(256,N))
        a=km.fit_predict(X)
    else: raise ValueError(policy)
    return a

def sparse_bytes(M):
    M=M.tocsr()
    return int(M.data.nbytes+M.indices.nbytes+M.indptr.nbytes)

rows=[]; routing_rows=[]; envelope_rows=[]
all_assign={}
for k in KS:
  for policy in POLICIES:
    a=assignments(policy,k); all_assign[(policy,k)]=a
    counts=np.bincount(a,minlength=k); bload=np.bincount(a,weights=bytes_arr,minlength=k)
    cut=sum(1 for i,j in knn_edges if a[i]!=a[j])/len(knn_edges)
    # shard centroids and exact nonnegative envelopes
    cent=[]; env=[]; env_bytes=0; centroid_bytes=0
    shard_idx=[]
    for s in range(k):
        idx=np.where(a==s)[0]; shard_idx.append(idx)
        if len(idx)==0:
            c=sparse.csr_matrix((1,V)); e=sparse.csr_matrix((1,V))
        else:
            c=X[idx].mean(axis=0)
            c=sparse.csr_matrix(c); c=normalize(c)
            # sparse max by column
            e=sparse.csr_matrix(X[idx].max(axis=0))
        cent.append(c); env.append(e)
        centroid_bytes+=sparse_bytes(c); env_bytes+=sparse_bytes(e)
    C=sparse.vstack(cent).tocsr(); E=sparse.vstack(env).tocsr()
    rows.append({
        'policy':policy,'shards':k,'chunks':N,'raw_body_bytes':raw_bytes,
        'chunk_count_cv':float(counts.std()/counts.mean()),
        'byte_load_cv':float(bload.std()/bload.mean()),
        'max_byte_share':float(bload.max()/bload.sum()),
        'semantic_knn_edge_cut':float(cut),
        'centroid_summary_bytes':centroid_bytes,'envelope_summary_bytes':env_bytes,
        'centroid_to_body_ratio':centroid_bytes/raw_bytes,'envelope_to_body_ratio':env_bytes/raw_bytes
    })
    # Approximate centroid routing top-m
    cs=(Q @ C.T).toarray()
    for m0 in [1,2,4,8]:
        m=min(m0,k)
        recalls=[]; exactsets=[]; scored=[]; span=[]
        for qi in range(Qn):
            ss=topk_from_scores(cs[qi],m)
            cand=np.concatenate([shard_idx[s] for s in ss]) if len(ss) else np.array([],dtype=int)
            sc=central_scores[qi,cand]
            local=cand[topk_from_scores(sc,TOPK)] if len(cand) else np.array([],dtype=int)
            ref=set(map(int,central_top[qi])); got=set(map(int,local))
            recalls.append(len(ref & got)/TOPK)
            exactsets.append(ref==got)
            scored.append(len(cand)/N)
            span.append(len(set(a[list(ref)])))
        routing_rows.append({
            'policy':policy,'shards':k,'router':f'centroid_top{m}',
            'mean_recall_at5':float(np.mean(recalls)),
            'exact_top5_set_rate':float(np.mean(exactsets)),
            'mean_scored_fraction':float(np.mean(scored)),
            'mean_shards_contacted':float(m),
            'mean_reference_shard_span':float(np.mean(span))
        })
    # Exact envelope branch-and-bound
    ub=(Q @ E.T).toarray()
    contacts=[]; scoredfr=[]; exact=[]
    for qi in range(Qn):
        order=np.argsort(-ub[qi])
        best_idx=np.array([],dtype=int); best_sc=np.array([],dtype=float)
        ncontact=0; nscored=0
        for pos,s in enumerate(order):
            idx=shard_idx[int(s)]; ncontact+=1; nscored+=len(idx)
            vals=central_scores[qi,idx]
            if len(best_idx)==0:
                best_idx=idx.copy(); best_sc=vals.copy()
            else:
                best_idx=np.concatenate([best_idx,idx]); best_sc=np.concatenate([best_sc,vals])
            take=topk_from_scores(best_sc,TOPK)
            best_idx=best_idx[take]; best_sc=best_sc[take]
            tau=float(best_sc.min()) if len(best_sc)>=TOPK else -1
            nextub=float(ub[qi,order[pos+1]]) if pos+1<len(order) else -1
            if len(best_sc)>=TOPK and tau+1e-12>=nextub:
                break
        ref=set(map(int,central_top[qi])); got=set(map(int,best_idx))
        contacts.append(ncontact); scoredfr.append(nscored/N); exact.append(ref==got)
    envelope_rows.append({
        'policy':policy,'shards':k,'router':'exact_envelope_bnb',
        'exact_top5_set_rate':float(np.mean(exact)),
        'mean_shards_contacted':float(np.mean(contacts)),
        'median_shards_contacted':float(np.median(contacts)),
        'p90_shards_contacted':float(np.quantile(contacts,.9)),
        'mean_shard_fraction_contacted':float(np.mean(contacts)/k),
        'mean_scored_fraction':float(np.mean(scoredfr))
    })

pd.DataFrame(rows).to_csv(OUT/'do026_sharding.csv',index=False)
pd.DataFrame(routing_rows).to_csv(OUT/'do026_centroid_routing.csv',index=False)
pd.DataFrame(envelope_rows).to_csv(OUT/'do026_envelope_routing.csv',index=False)
pd.DataFrame([{'chunks':N,'sections':sum(c['kind']=='section' for c in chunks),'tables':sum(c['kind']=='table' for c in chunks),'raw_body_bytes':raw_bytes,'tfidf_features':V,'tfidf_nnz':X.nnz,'tfidf_sparse_bytes':sparse_bytes(X),'queries':Qn,'semantic_knn_edges':len(knn_edges)}]).to_csv(OUT/'do026_workload_summary.csv',index=False)

# ---------- actual worker-process batching benchmark ----------
# Use semantic 8-shard layout; compare centralized vs all-shard distributed reduction.
A=all_assign[('semantic_kmeans',8)]
shards=[np.where(A==s)[0] for s in range(8)]
# repeat the real query set only to amortize process startup; no new query content is introduced.
reps=max(1, math.ceil(4096/Qn)); Qbench=sparse.vstack([Q]*reps)[:4096].tocsr()

def worker_loop(Xsub, global_ids, iq, oq):
    while True:
        msg=iq.get()
        if msg is None: break
        bid, qb=msg
        M=(qb @ Xsub.T).toarray()
        out=[]
        for r in range(M.shape[0]):
            loc=topk_from_scores(M[r],TOPK)
            out.append((global_ids[loc].tolist(), M[r,loc].tolist()))
        oq.put((bid,out))

def distributed_broadcast(batch_size):
    ctx=mp.get_context('fork')
    inqs=[]; outq=ctx.Queue(); procs=[]
    for idx in shards:
        iq=ctx.Queue(maxsize=4); p=ctx.Process(target=worker_loop,args=(X[idx],idx,iq,outq)); p.start(); inqs.append(iq); procs.append(p)
    t0=time.perf_counter(); bid=0; nb=0
    for st in range(0,Qbench.shape[0],batch_size):
        qb=Qbench[st:st+batch_size]
        for iq in inqs: iq.put((bid,qb))
        bid+=1; nb+=1
    # collect and merge every shard batch response; merge cost is actually executed.
    received=defaultdict(list)
    for _ in range(nb*8):
        b,res=outq.get(); received[b].append(res)
    checksum=0
    for b,lists in received.items():
        bs=len(lists[0])
        for r in range(bs):
            ids0=[]; sc0=[]
            for shardres in lists:
                ii,ss=shardres[r]; ids0.extend(ii); sc0.extend(ss)
            sca=np.array(sc0); ida=np.array(ids0)
            take=topk_from_scores(sca,TOPK); checksum+=int(ida[take].sum())
    wall=time.perf_counter()-t0
    for iq in inqs: iq.put(None)
    for p in procs: p.join()
    return wall,checksum

bench=[]
# centralized sparse batch baseline
for bs in [1,8,32,128]:
    t0=time.perf_counter(); checksum=0
    for st in range(0,Qbench.shape[0],bs):
        M=(Qbench[st:st+bs] @ X.T).toarray()
        for r in range(M.shape[0]): checksum+=int(topk_from_scores(M[r],TOPK).sum())
    bench.append({'mode':'centralized','shards':1,'batch_size':bs,'queries':Qbench.shape[0],'wall_seconds':time.perf_counter()-t0,'checksum':checksum})
for bs in [1,8,32,128]:
    wall,checksum=distributed_broadcast(bs)
    bench.append({'mode':'distributed_broadcast','shards':8,'batch_size':bs,'queries':Qbench.shape[0],'wall_seconds':wall,'checksum':checksum})
benchdf=pd.DataFrame(bench); benchdf['queries_per_second']=benchdf['queries']/benchdf['wall_seconds']
benchdf.to_csv(OUT/'do026_process_benchmark.csv',index=False)

# Save metadata and query definitions for evidence.
with open(OUT/'do026_queries.json','w',encoding='utf-8') as f: json.dump({'queries':queries,'meta':qmeta},f,ensure_ascii=False,indent=2)
with open(OUT/'do026_chunk_inventory.json','w',encoding='utf-8') as f:
    json.dump([{k:c[k] for k in ['id','kind','heading','chapter','utf8_bytes','chars']} for c in chunks],f,ensure_ascii=False,indent=2)

summary={
 'workload':pd.read_csv(OUT/'do026_workload_summary.csv').iloc[0].to_dict(),
 'best_sharding_by_cut':pd.read_csv(OUT/'do026_sharding.csv').sort_values(['shards','semantic_knn_edge_cut']).groupby('shards').first().reset_index().to_dict('records'),
 'envelope':pd.read_csv(OUT/'do026_envelope_routing.csv').to_dict('records'),
 'bench':benchdf.to_dict('records')
}
with open(OUT/'do026_summary.json','w') as f: json.dump(summary,f,indent=2)
print(json.dumps(summary,indent=2)[:16000])
